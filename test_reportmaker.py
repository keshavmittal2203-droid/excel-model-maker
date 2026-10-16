"""Report-builder tests: generate a sample workbook, validate it, exercise the CLI.

These tests prove structure and formatting of the generated file. They are not
Excel certification: opening the output in Excel remains the final check.
"""
from __future__ import annotations

from datetime import date, datetime
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
import zipfile

from openpyxl import load_workbook
from openpyxl.chart import BarChart

from reportmaker import ReportError, build_workbook, parse_spec, write_report
from reportmaker.spec import spec_from_file

ROOT = Path(__file__).resolve().parent


def sample_spec(**overrides):
    spec = {
        "title": "Test Report",
        "sheet_name": "Test",
        "columns": [
            {"key": "month", "label": "Month", "type": "date"},
            {"key": "units", "label": "Units Sold", "type": "integer", "total": True},
            {"key": "revenue", "label": "Revenue", "type": "currency", "total": True},
            {"key": "growth", "label": "Growth", "type": "percent", "conditional": "color_scale"},
            {"key": "margin", "label": "Margin", "type": "percent", "conditional": "data_bar"},
        ],
        "rows": [
            {"month": "2026-01-01", "units": 100, "revenue": 1000.5, "growth": 0.05, "margin": 0.30},
            {"month": "2026-02-01", "units": 110, "revenue": 1210.0, "growth": -0.02, "margin": 0.32},
            {"month": "2026-03-01", "units": 90, "revenue": 990.25, "growth": 0.08, "margin": 0.28},
        ],
        "chart": {"type": "bar", "title": "Revenue by Month", "categories": "month", "series": ["revenue"]},
    }
    spec.update(overrides)
    return spec


def as_date(value):
    if isinstance(value, datetime):
        return value.date()
    return value


class BuilderTests(unittest.TestCase):
    def test_generates_styled_sample_workbook(self):
        spec = parse_spec(sample_spec())
        worksheet = build_workbook(spec).active

        self.assertEqual(worksheet.title, "Test")
        self.assertEqual(worksheet["A1"].value, "Test Report")
        self.assertIn("A1:E1", [str(rng) for rng in worksheet.merged_cells.ranges])

        header = worksheet["A2"]
        self.assertEqual(header.value, "Month")
        self.assertTrue(header.font.bold)
        self.assertTrue(str(header.font.color.rgb).upper().endswith("FFFFFF"))
        self.assertTrue(str(header.fill.fgColor.rgb).upper().endswith("1F4E79"))
        self.assertEqual(worksheet.freeze_panes, "A3")

        self.assertEqual(worksheet["A3"].value, date(2026, 1, 1))
        self.assertEqual(worksheet["A3"].number_format, "yyyy-mm-dd")
        self.assertEqual(worksheet["B3"].number_format, "#,##0")
        self.assertEqual(worksheet["C3"].number_format, '"$"#,##0.00')
        self.assertEqual(worksheet["D3"].number_format, "0.0%")
        self.assertEqual(worksheet["B3"].alignment.horizontal, "right")
        self.assertEqual(worksheet["A3"].alignment.horizontal, "left")

        self.assertTrue(str(worksheet["A4"].fill.fgColor.rgb).upper().endswith("EAF1F8"))
        self.assertIsNone(worksheet["A3"].fill.fill_type)
        self.assertIsNone(worksheet["A5"].fill.fill_type)

        self.assertEqual(worksheet.auto_filter.ref, "A2:E5")

        self.assertEqual(worksheet["A6"].value, "Total")
        self.assertEqual(worksheet["B6"].value, "=SUM(B3:B5)")
        self.assertEqual(worksheet["C6"].value, "=SUM(C3:C5)")
        self.assertIsNone(worksheet["D6"].value)
        self.assertTrue(worksheet["B6"].font.bold)
        self.assertEqual(worksheet["B6"].border.top.style, "double")

        rules = {str(cf.sqref): [rule.type for rule in cf.rules] for cf in worksheet.conditional_formatting}
        self.assertEqual(rules.get("D3:D5"), ["colorScale"])
        self.assertEqual(rules.get("E3:E5"), ["dataBar"])

        self.assertEqual(worksheet.column_dimensions["A"].width, 12.0)
        self.assertGreaterEqual(worksheet.column_dimensions["C"].width, 10.0)

        self.assertEqual(len(worksheet._charts), 1)
        self.assertIsInstance(worksheet._charts[0], BarChart)
        self.assertEqual(len(worksheet._charts[0].series), 1)

        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "sample.xlsx"
            manifest = write_report(spec, output)
            self.assertEqual(manifest["status"], "generated")
            self.assertFalse(manifest["ready"])
            self.assertEqual(manifest["data_rows"], 3)
            self.assertEqual(manifest["total_columns"], ["units", "revenue"])
            self.assertEqual(manifest["conditional_format_columns"], ["growth", "margin"])
            self.assertEqual(manifest["chart"], "bar")
            self.assertEqual(len(manifest["output_sha256"]), 64)
            self.assertGreater(output.stat().st_size, 2000)

            with self.assertRaises(ReportError):
                write_report(spec, output)
            write_report(spec, output, overwrite=True)

            with zipfile.ZipFile(output) as archive:
                chart_parts = [name for name in archive.namelist() if name.startswith("xl/charts/")]
                self.assertTrue(chart_parts)
                self.assertIn(b"Revenue by Month", archive.read(chart_parts[0]))

            reopened = load_workbook(output).active
            self.assertEqual(reopened.freeze_panes, "A3")
            self.assertEqual(reopened["B6"].value, "=SUM(B3:B5)")
            self.assertEqual(reopened["C3"].number_format, '"$"#,##0.00')
            self.assertEqual(as_date(reopened["A3"].value), date(2026, 1, 1))
            self.assertEqual(len(reopened._charts), 1)

    def test_layout_without_title_has_header_on_row_one(self):
        spec = parse_spec(sample_spec(title=None, sheet_name="Plain", chart=None))
        worksheet = build_workbook(spec).active
        self.assertEqual(worksheet["A1"].value, "Month")
        self.assertEqual(worksheet.freeze_panes, "A2")
        self.assertEqual(list(worksheet.merged_cells.ranges), [])
        self.assertEqual(worksheet.auto_filter.ref, "A1:E4")
        self.assertEqual(worksheet["B5"].value, "=SUM(B2:B4)")
        self.assertEqual(len(worksheet._charts), 0)


class SpecValidationTests(unittest.TestCase):
    def assert_rejects(self, spec, fragment):
        with self.assertRaises(ReportError) as context:
            parse_spec(spec)
        self.assertIn(fragment, str(context.exception))

    def test_column_and_sheet_validation(self):
        spec = sample_spec()
        spec["tital"] = "typo"
        self.assert_rejects(spec, "unknown field")
        spec = sample_spec()
        del spec["columns"]
        self.assert_rejects(spec, "missing required field")
        self.assert_rejects(sample_spec(columns=[]), "columns must be a non-empty")
        self.assert_rejects(sample_spec(columns=[{"label": "No key"}]), "missing required field(s): key")
        self.assert_rejects(sample_spec(columns=[{"key": "x", "type": "money"}]), "type must be one of")
        self.assert_rejects(sample_spec(columns=[{"key": "x", "type": "number", "decimals": -1}]), "decimals must be an integer")
        self.assert_rejects(sample_spec(columns=[{"key": "x", "type": "number", "decimals": "2"}]), "decimals must be an integer")
        self.assert_rejects(sample_spec(columns=[{"key": "x", "currency_symbol": "$"}]), "only valid for currency")
        self.assert_rejects(sample_spec(columns=[{"key": "x", "type": "text", "total": True}]), "total is supported only")
        self.assert_rejects(sample_spec(columns=[{"key": "x", "type": "percent", "total": True}]), "total is supported only")
        self.assert_rejects(sample_spec(columns=[{"key": "x", "type": "date", "conditional": "data_bar"}]), "supported only for numeric")
        self.assert_rejects(sample_spec(columns=[{"key": "x", "type": "number", "conditional": "heat_map"}]), "conditional.type must be one of")
        duplicate = [{"key": "x"}, {"key": "x"}]
        self.assert_rejects(sample_spec(columns=duplicate), "Duplicate column key")
        self.assert_rejects(sample_spec(sheet_name="x" * 32), "31 characters")
        self.assert_rejects(sample_spec(sheet_name="Bad:Name"), "must not contain")
        self.assert_rejects(sample_spec(sheet_name="'Quoted'"), "apostrophe")

    def test_row_validation(self):
        rows = [dict(row) for row in sample_spec()["rows"]]
        missing = [dict(row) for row in rows]
        del missing[1]["units"]
        self.assert_rejects(sample_spec(rows=missing), "missing units")
        extra = [dict(row) for row in rows]
        extra[0]["profit"] = 1
        self.assert_rejects(sample_spec(rows=extra), "unknown profit")
        self.assert_rejects(sample_spec(rows=[]), "rows must be a non-empty")
        bad = [dict(row) for row in rows]
        bad[0]["units"] = True
        self.assert_rejects(sample_spec(rows=bad), "must be a JSON number")
        bad = [dict(row) for row in rows]
        bad[0]["revenue"] = "lots"
        self.assert_rejects(sample_spec(rows=bad), "must be a JSON number")
        bad = [dict(row) for row in rows]
        bad[0]["units"] = 1.5
        self.assert_rejects(sample_spec(rows=bad), "whole number")
        bad = [dict(row) for row in rows]
        bad[0]["month"] = "2026-02-30"
        self.assert_rejects(sample_spec(rows=bad), "not a valid calendar date")
        bad = [dict(row) for row in rows]
        bad[0]["month"] = "01/02/2026"
        self.assert_rejects(sample_spec(rows=bad), "must be an ISO date")

    def test_chart_validation(self):
        base = sample_spec()["chart"]
        self.assert_rejects(sample_spec(chart={**base, "type": "scatter"}), "chart.type must be one of")
        self.assert_rejects(sample_spec(chart={**base, "series": ["profit"]}), "unknown column")
        self.assert_rejects(sample_spec(chart={**base, "series": ["month"]}), "must be a numeric column")
        self.assert_rejects(sample_spec(chart={**base, "categories": "revenue"}), "must be a text or date column")
        self.assert_rejects(sample_spec(chart={**base, "series": ["revenue", "revenue"]}), "duplicate column keys")
        self.assert_rejects(sample_spec(chart={**base, "type": "pie", "series": ["revenue", "units"]}), "exactly one series")
        self.assert_rejects(sample_spec(chart={**base, "position": "ZZ"}), "cell address")

    def test_file_level_rejections(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            duplicate = root / "duplicate.json"
            duplicate.write_text('{"columns": [], "columns": []}', encoding="utf-8")
            with self.assertRaises(ReportError) as context:
                spec_from_file(duplicate)
            self.assertIn("Duplicate JSON key", str(context.exception))
            constant = root / "constant.json"
            constant.write_text('{"columns": [], "rows": [], "x": NaN}', encoding="utf-8")
            with self.assertRaises(ReportError) as context:
                spec_from_file(constant)
            self.assertIn("Invalid JSON numeric constant", str(context.exception))
            malformed = root / "malformed.json"
            malformed.write_text('{"columns": ', encoding="utf-8")
            with self.assertRaises(ReportError) as context:
                spec_from_file(malformed)
            self.assertIn("Invalid JSON", str(context.exception))
            array = root / "array.json"
            array.write_text("[]", encoding="utf-8")
            with self.assertRaises(ReportError) as context:
                spec_from_file(array)
            self.assertIn("must be a JSON object", str(context.exception))
        with self.assertRaises(ReportError) as context:
            spec_from_file(Path(tempfile.gettempdir()) / "definitely-missing-report.json")
        self.assertIn("not found", str(context.exception))


class CliTests(unittest.TestCase):
    def test_cli_end_to_end_on_shipped_example(self):
        cli = ROOT / "cli.py"
        example = ROOT / "report.example.json"
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "report.xlsx"
            command = [sys.executable, str(cli), "--input", str(example), "--output", str(output)]
            first = subprocess.run(command, capture_output=True, text=True, encoding="utf-8", cwd=str(ROOT))
            self.assertEqual(first.returncode, 0, first.stdout + first.stderr)
            payload = json.loads(first.stdout)
            self.assertEqual(payload["status"], "generated")
            self.assertEqual(payload["chart"], "bar")
            self.assertTrue(output.is_file())
            self.assertEqual(load_workbook(output).active["A2"].value, "Month")

            before = output.read_bytes()
            second = subprocess.run(command, capture_output=True, text=True, encoding="utf-8", cwd=str(ROOT))
            self.assertEqual(second.returncode, 2)
            self.assertIn("already exists", json.loads(second.stdout)["error"])
            self.assertEqual(output.read_bytes(), before)

            forced = subprocess.run(command + ["--force"], capture_output=True, text=True, encoding="utf-8", cwd=str(ROOT))
            self.assertEqual(forced.returncode, 0, forced.stdout + forced.stderr)

            missing = subprocess.run(
                [sys.executable, str(cli), "--input", str(Path(directory) / "nope.json"), "--output", str(Path(directory) / "x.xlsx")],
                capture_output=True, text=True, encoding="utf-8", cwd=str(ROOT),
            )
            self.assertEqual(missing.returncode, 2)
            self.assertIn("not found", json.loads(missing.stdout)["error"])


if __name__ == "__main__":
    unittest.main()
