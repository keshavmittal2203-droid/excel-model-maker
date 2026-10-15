"""Assemble a styled presentation workbook from a validated report spec.

Styling: colored bold header, frozen panes, banded data rows, per-column
number formats, estimated column widths, SUM totals as real Excel formulas,
conditional formatting, and an embedded chart.

Honest limits: this writes a new workbook; it never edits templates, never
evaluates formulas, and never caches results. Excel computes the SUM totals
when the file is opened.
"""
from __future__ import annotations

import hashlib
import io
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from . import theme
from .errors import ReportError
from .spec import Column, ReportSpec

MIN_WIDTH = 10.0
MAX_WIDTH = 60.0
WIDTH_PADDING = 2.0
CHART_WIDTH_CM = 20.0
CHART_HEIGHT_CM = 10.0
CHART_GAP_ROWS = 2


def number_format(column: Column) -> str:
    if column.type == "text":
        return "General"
    if column.type == "date":
        return "yyyy-mm-dd"
    digits = "#,##0" + ("." + "0" * column.decimals if column.decimals else "")
    if column.type == "currency":
        return f'"{column.currency_symbol}"{digits}'
    if column.type == "percent":
        return "0" + ("." + "0" * column.decimals if column.decimals else "") + "%"
    return digits


def display_text(column: Column, value: Any) -> str:
    """Approximate how Excel will render a value; used for width estimation."""
    if value is None:
        return ""
    if column.type == "date":
        return value.strftime("%Y-%m-%d")
    if column.type == "percent":
        return f"{value * 100:,.{column.decimals}f}%"
    if column.type == "currency":
        return f"{column.currency_symbol}{value:,.{column.decimals}f}"
    if column.type in ("integer", "number"):
        return f"{value:,.{column.decimals}f}"
    return str(value)


def _estimated_widths(spec: ReportSpec) -> list[float]:
    widths = []
    for column in spec.columns:
        lengths = [len(column.label)]
        for row in spec.rows:
            lengths.append(len(display_text(column, row[column.key])))
        if column.total:
            total = sum(row[column.key] for row in spec.rows)
            lengths.append(len(display_text(column, total)))
        widths.append(min(MAX_WIDTH, max(MIN_WIDTH, max(lengths) + WIDTH_PADDING)))
    return widths


def _chart_object(chart_type: str):
    if chart_type == "bar":
        chart = BarChart()
        chart.type = "col"
        return chart
    if chart_type == "line":
        return LineChart()
    return PieChart()


def _add_chart(worksheet: Worksheet, spec: ReportSpec, header_row: int, first_data_row: int, last_data_row: int, anchor_row: int) -> None:
    chart_spec = spec.chart
    assert chart_spec is not None
    by_key = {column.key: (index, column) for index, column in enumerate(spec.columns, start=1)}
    chart = _chart_object(chart_spec.type)
    for key in chart_spec.series:
        index, _ = by_key[key]
        chart.add_data(
            Reference(worksheet, min_col=index, min_row=header_row, max_row=last_data_row),
            titles_from_data=True,
        )
    category_index, category_column = by_key[chart_spec.categories]
    chart.set_categories(Reference(worksheet, min_col=category_index, min_row=first_data_row, max_row=last_data_row))
    if chart_spec.title is not None:
        chart.title = chart_spec.title
    else:
        series_labels = ", ".join(by_key[key][1].label for key in chart_spec.series)
        chart.title = f"{series_labels} by {category_column.label}"
    if isinstance(chart, PieChart):
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showPercent = True
    else:
        chart.x_axis.title = category_column.label
        if len(chart_spec.series) == 1:
            chart.y_axis.title = by_key[chart_spec.series[0]][1].label
    chart.width = CHART_WIDTH_CM
    chart.height = CHART_HEIGHT_CM
    worksheet.add_chart(chart, chart_spec.position or f"A{anchor_row}")


def build_workbook(spec: ReportSpec) -> Workbook:
    """Build the styled workbook in memory (no formulas evaluated)."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = spec.sheet_name
    workbook.properties.title = spec.title or spec.sheet_name
    workbook.properties.creator = "excel-model-maker report builder"

    ncols = len(spec.columns)
    last_col = get_column_letter(ncols)

    header_row = 1
    if spec.title:
        header_row = 2
        worksheet.merge_cells(f"A1:{last_col}1")
        title_cell = worksheet.cell(row=1, column=1, value=spec.title)
        title_cell.font = theme.TITLE_FONT
        title_cell.alignment = theme.LEFT_ALIGN
        worksheet.row_dimensions[1].height = theme.TITLE_ROW_HEIGHT

    for index, column in enumerate(spec.columns, start=1):
        cell = worksheet.cell(row=header_row, column=index, value=column.label)
        cell.font = theme.HEADER_FONT
        cell.fill = theme.HEADER_FILL
        cell.alignment = theme.HEADER_ALIGN
    worksheet.row_dimensions[header_row].height = theme.HEADER_ROW_HEIGHT
    worksheet.freeze_panes = f"A{header_row + 1}"

    first_data_row = header_row + 1
    last_data_row = header_row + len(spec.rows)
    for row_index, row in enumerate(spec.rows, start=first_data_row):
        banded = (row_index - first_data_row) % 2 == 1
        for index, column in enumerate(spec.columns, start=1):
            cell = worksheet.cell(row=row_index, column=index, value=row[column.key])
            cell.font = theme.BODY_FONT
            cell.number_format = number_format(column)
            cell.alignment = theme.LEFT_ALIGN if column.type in ("text", "date") else theme.RIGHT_ALIGN
            if banded:
                cell.fill = theme.BAND_FILL

    worksheet.auto_filter.ref = f"A{header_row}:{last_col}{last_data_row}"

    total_row = None
    if any(column.total for column in spec.columns):
        total_row = last_data_row + 1
        for index, column in enumerate(spec.columns, start=1):
            cell = worksheet.cell(row=total_row, column=index)
            if index == 1:
                # The label occupies the first column; a total on that column is skipped by design.
                cell.value = "Total"
            elif column.total:
                letter = get_column_letter(index)
                cell.value = f"=SUM({letter}{first_data_row}:{letter}{last_data_row})"
            cell.font = theme.TOTAL_FONT
            cell.fill = theme.TOTAL_FILL
            cell.border = theme.TOTAL_TOP_BORDER
            cell.number_format = number_format(column)
            cell.alignment = theme.LEFT_ALIGN if column.type in ("text", "date") else theme.RIGHT_ALIGN

    for index, column in enumerate(spec.columns, start=1):
        if column.conditional is None:
            continue
        letter = get_column_letter(index)
        cell_range = f"{letter}{first_data_row}:{letter}{last_data_row}"
        if column.conditional == "color_scale":
            worksheet.conditional_formatting.add(
                cell_range,
                ColorScaleRule(start_type="min", start_color=theme.SCALE_LOW, mid_type="percentile", mid_value=50, mid_color=theme.SCALE_MID, end_type="max", end_color=theme.SCALE_HIGH),
            )
        else:
            worksheet.conditional_formatting.add(
                cell_range,
                DataBarRule(start_type="min", end_type="max", color=theme.DATA_BAR_COLOR, showValue=True),
            )

    for index, width in enumerate(_estimated_widths(spec), start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width

    if spec.chart is not None:
        anchor_row = (total_row if total_row is not None else last_data_row) + CHART_GAP_ROWS
        _add_chart(worksheet, spec, header_row, first_data_row, last_data_row, anchor_row)

    return workbook


def write_report(spec: ReportSpec, output: str | Path, *, overwrite: bool = False) -> dict[str, Any]:
    """Build, save (exclusively unless overwrite), verify, and return a manifest."""
    target = Path(output)
    if target.parent and not target.parent.exists():
        raise ReportError(f"Output directory does not exist: {target.parent}")
    workbook = build_workbook(spec)
    buffer = io.BytesIO()
    workbook.save(buffer)
    payload = buffer.getvalue()
    try:
        with target.open("wb" if overwrite else "xb") as destination:
            destination.write(payload)
    except FileExistsError as exc:
        raise ReportError(f"Output already exists: {target} (use --force to overwrite)") from exc
    except OSError as exc:
        raise ReportError(f"Could not write output {target}: {exc}") from exc
    try:
        load_workbook(io.BytesIO(payload))
    except Exception as exc:  # openpyxl raises several exception types on bad packages
        raise ReportError(f"Generated workbook failed verification: {exc}") from exc
    return {
        "status": "generated",
        "ready": False,
        "note": "Presentation report only: Excel evaluates the SUM formulas when the file is opened.",
        "output": str(target),
        "output_sha256": hashlib.sha256(payload).hexdigest(),
        "sheet_name": spec.sheet_name,
        "columns": [column.key for column in spec.columns],
        "data_rows": len(spec.rows),
        "total_columns": [column.key for column in spec.columns if column.total],
        "conditional_format_columns": [column.key for column in spec.columns if column.conditional],
        "chart": spec.chart.type if spec.chart else None,
    }
