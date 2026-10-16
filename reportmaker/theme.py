"""Central styling constants for generated reports."""
from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

BRAND_DARK = "1F4E79"
BAND_COLOR = "EAF1F8"
TOTAL_COLOR = "D9E2F3"

TITLE_FONT = Font(name="Calibri", size=14, bold=True, color=BRAND_DARK)
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Calibri", size=11)
TOTAL_FONT = Font(name="Calibri", size=11, bold=True)

HEADER_FILL = PatternFill("solid", fgColor=BRAND_DARK)
BAND_FILL = PatternFill("solid", fgColor=BAND_COLOR)
TOTAL_FILL = PatternFill("solid", fgColor=TOTAL_COLOR)

HEADER_ALIGN = Alignment(horizontal="center", vertical="center")
LEFT_ALIGN = Alignment(horizontal="left", vertical="center")
RIGHT_ALIGN = Alignment(horizontal="right", vertical="center")

TOTAL_TOP_BORDER = Border(top=Side(style="double", color=BRAND_DARK))

TITLE_ROW_HEIGHT = 22.0
HEADER_ROW_HEIGHT = 20.0

# Conditional-formatting palette (Excel's classic three-color scale).
SCALE_LOW = "F8696B"
SCALE_MID = "FFEB84"
SCALE_HIGH = "63BE7B"
DATA_BAR_COLOR = "638EC6"
