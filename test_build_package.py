"""Unit and regression tests for build_package.py orchestrator."""
from __future__ import annotations

import copy
import json
from pathlib import Path
import tempfile
import unittest

from build_package import (
    build_package,
    check_com_capabilities,
    compute_file_sha,
    verify_output_directory,
)
from modelmaker import sha

ROOT = Path(__file__).resolve().parent


class BuildPackageTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.test_dir = Path(self.temp_dir.name)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_forbidden_directories_rejected(self):
        # Repositories roots and protected source trees must be rejected
        with self.assertRaises(RuntimeError):
            verify_output_directory(ROOT)
        with self.assertRaises(RuntimeError):
            verify_output_directory(ROOT / "originals")
        with self.assertRaises(RuntimeError):
            verify_output_directory(ROOT / "proposals")
        with self.assertRaises(RuntimeError):
            verify_output_directory(ROOT / "docs")
        with self.assertRaises(RuntimeError):
            verify_output_directory(ROOT / "work")
        with self.assertRaises(RuntimeError):
            verify_output_directory(ROOT / ".git")

    def test_overwrite_rejected_without_force(self):
        out_dir = self.test_dir / "output"
        out_dir.mkdir(parents=True, exist_ok=True)
        # Create an existing file inside output directory
        (out_dir / "dummy.txt").write_text("existing content", encoding="utf-8")

        req_path = ROOT / "examples" / "marksans-package.json"
        with self.assertRaises(RuntimeError) as cm:
            build_package(
                request_path=req_path,
                output_dir=out_dir,
                allow_proposal=True,
                force=False,
                skip_com=True,
            )
        self.assertIn("already exists and is not empty", str(cm.exception))
        self.assertIn("--force", str(cm.exception))

    def test_force_rejected_on_non_package_arbitrary_dir(self):
        # If directory is not a known run dir and lacks package.manifest.json, force must fail
        # Create a mock arbitrary dir outside temp
        arbitrary_dir = ROOT / "test_scratch_nonpackage"
        try:
            arbitrary_dir.mkdir(parents=True, exist_ok=True)
            (arbitrary_dir / "important_file.txt").write_text("critical data", encoding="utf-8")
            with self.assertRaises(RuntimeError) as cm:
                verify_output_directory(arbitrary_dir, force=True)
            self.assertTrue(
                "does not contain a package.manifest.json" in str(cm.exception)
                or "fresh output directory" in str(cm.exception)
            )
        finally:
            import shutil
            shutil.rmtree(arbitrary_dir, ignore_errors=True)

    def test_com_capabilities_structure(self):
        caps = check_com_capabilities()
        self.assertIn("excel_com", caps)
        self.assertIn("platform", caps)
        self.assertIsInstance(caps["excel_com"], bool)

    def test_source_sha_mismatch_rejected(self):
        out_dir = self.test_dir / "output"
        req_text = (ROOT / "examples" / "marksans-package.json").read_text(encoding="utf-8")
        req = json.loads(req_text)
        # Corrupt expected source SHA
        req["workbook"]["source_sha256"] = "0" * 64
        corrupted_req = self.test_dir / "bad_req.json"
        corrupted_req.write_text(json.dumps(req), encoding="utf-8")

        with self.assertRaises(ValueError) as cm:
            build_package(
                request_path=corrupted_req,
                output_dir=out_dir,
                allow_proposal=True,
                force=True,
                skip_com=True,
            )
        self.assertIn("Workbook source SHA mismatch", str(cm.exception))

    def test_unapproved_pack_rejected_without_allow_proposal(self):
        out_dir = self.test_dir / "output"
        req_path = ROOT / "examples" / "marksans-package.json"

        with self.assertRaises(Exception) as cm:
            build_package(
                request_path=req_path,
                output_dir=out_dir,
                allow_proposal=False,
                force=True,
                skip_com=True,
            )
        err_msg = str(cm.exception).lower()
        self.assertTrue("proposal" in err_msg or "approval" in err_msg or "blocked" in err_msg)

    def test_build_package_baseline_end_to_end_skip_com(self):
        out_dir = self.test_dir / "output_baseline"
        req_path = ROOT / "examples" / "marksans-package.json"

        manifest = build_package(
            request_path=req_path,
            output_dir=out_dir,
            allow_proposal=True,
            force=True,
            skip_com=True,
        )

        self.assertEqual(manifest["package_id"], "marksans-curated-overview-2025")
        self.assertEqual(manifest["status"], "validating")
        self.assertFalse(manifest["ready"])
        self.assertIn("human_pack_approval", manifest["pending_gates"])
        self.assertIn("render_review", manifest["pending_gates"])
        self.assertIn("com_capabilities", manifest)

        # Check output artifacts
        wb_file = out_dir / "Marksans-Recalculated.xlsx"
        manifest_file = out_dir / "package.manifest.json"

        self.assertTrue(wb_file.exists())
        self.assertTrue(manifest_file.exists())

        # Verify no self-referential manifest hash
        artifact_files = [a["file"] for a in manifest["output_artifacts"]]
        self.assertNotIn("package.manifest.json", artifact_files)

        # Check reconciliation table entries
        reconciled_metrics = {m["metric"]: m for m in manifest["reconciliation_summary"]}
        self.assertEqual(reconciled_metrics["Risk-Free Rate (Rf)"]["displayed"], "7.30%")
        self.assertEqual(reconciled_metrics["Pre-tax Cost of Debt"]["displayed"], "8.00%")
        self.assertEqual(reconciled_metrics["Post-tax Cost of Debt (Kd)"]["displayed"], "5.60%")
        self.assertEqual(reconciled_metrics["Cost of Equity (Ke)"]["displayed"], "12.26%")
        self.assertEqual(reconciled_metrics["WACC (Target-Weighted)"]["displayed"], "12.11%")
        self.assertEqual(reconciled_metrics["FY25 Current Ratio"]["displayed"], "4.95")
        self.assertEqual(reconciled_metrics["FY25 Quick Ratio"]["displayed"], "3.07")

    def test_build_package_scenario_b_directional_effects(self):
        out_dir = self.test_dir / "output_scenario_b"
        req_path = ROOT / "examples" / "scenario-b-package.json"

        manifest = build_package(
            request_path=req_path,
            output_dir=out_dir,
            allow_proposal=True,
            force=True,
            skip_com=True,
        )

        self.assertEqual(manifest["package_id"], "marksans-scenario-b-stress-2025")
        self.assertFalse(manifest["ready"])

        reconciled_metrics = {m["metric"]: m for m in manifest["reconciliation_summary"]}
        # Directional verification: Rf increases to 8.00%
        self.assertEqual(reconciled_metrics["Risk-Free Rate (Rf)"]["displayed"], "8.00%")
        # Pre-tax Kd increases to 9.00%, post-tax Kd increases to 6.30%
        self.assertEqual(reconciled_metrics["Pre-tax Cost of Debt"]["displayed"], "9.00%")
        self.assertEqual(reconciled_metrics["Post-tax Cost of Debt (Kd)"]["displayed"], "6.30%")
        # Ke increases from 12.26% to 12.57%
        self.assertEqual(reconciled_metrics["Cost of Equity (Ke)"]["displayed"], "12.57%")
        # WACC increases from 12.11% to 12.43%
        self.assertEqual(reconciled_metrics["WACC (Target-Weighted)"]["displayed"], "12.43%")
        # Historical ratio metrics remain 100% identical
        self.assertEqual(reconciled_metrics["FY25 Current Ratio"]["displayed"], "4.95")
        self.assertEqual(reconciled_metrics["FY25 Quick Ratio"]["displayed"], "3.07")

    def test_workbook_inventory_matches_actual_sheets(self):
        """Regression test: verify active master template has 16 sheets (sectordata removed) and archived template has 17 sheets."""
        import openpyxl
        wb_path = ROOT / "originals" / "Marksans Keshav.xlsx"
        wb = openpyxl.load_workbook(str(wb_path), read_only=True)
        actual_sheets = wb.sheetnames
        expected_sheets = [
            "Model ->",
            "Main Financials",
            "Ratio Analysis",
            "Intrinsic Growth",
            "Comps",
            "WACC-DCF",
            "DCF 1st way",
            "DCF 2nd way",
            "Data ->",
            "compdata",
            "Data Sheet",
            "BS",
            "Beta",
            "RM",
            "peerdata",
            "Cash Flow",
        ]
        self.assertEqual(actual_sheets, expected_sheets)
        # Verify archived master retains original 17 sheets
        wb_arch = openpyxl.load_workbook(str(ROOT / "originals" / "archive" / "Marksans Keshav.xlsx"), read_only=True)
        self.assertEqual(wb_arch.sheetnames, expected_sheets + ["sectordata"])
        # Ensure none of the 7 hallucinated sheet names exist
        forbidden_hallucinated = [
            "PL",
            "common size statement",
            "historic forecast",
            "FCFF",
            "SOP",
            "cost of capital",
            "dividend discount model",
        ]
        for bad_sheet in forbidden_hallucinated:
            self.assertNotIn(bad_sheet, actual_sheets)

    def test_clean_checkout_guide_hashes_and_request_schema(self):
        """Regression test: verify AI_CLEAN_CHECKOUT_GUIDE.md has verified source hashes and valid request JSON."""
        import jsonschema
        guide_text = (ROOT / "docs" / "AI_CLEAN_CHECKOUT_GUIDE.md").read_text(encoding="utf-8")
        catalog = json.loads((ROOT / "source-catalog.json").read_text(encoding="utf-8"))
        for item in catalog["files"]:
            self.assertIn(
                item["sha256"],
                guide_text,
                f"Source hash for {item['file']} missing from AI_CLEAN_CHECKOUT_GUIDE.md",
            )

        # Verify example JSON request in guide validates against schema
        import re
        m = re.search(r"```json\s*(\{[\s\S]*?\})\s*```", guide_text)
        self.assertIsNotNone(m, "Could not find example JSON block in clean-checkout guide")
        example_json = json.loads(m.group(1))
        schema = json.loads((ROOT / "package-request.schema.json").read_text(encoding="utf-8"))
        # Should validate without raising jsonschema.ValidationError
        jsonschema.validate(instance=example_json, schema=schema)

    def test_reconciliation_summary_verified_identities_and_statuses(self):
        """Regression test: verify reconciliation summary evaluates mathematical identities and honest status labels."""
        from build_package import evaluate_reconciliation_summary
        wb_path = ROOT / "candidates" / "run_native_baseline_v5" / "Marksans-Recalculated.xlsx"
        if not wb_path.exists():
            wb_path = ROOT / "candidates" / "run_native_baseline_v4" / "Marksans-Recalculated.xlsx"
        if not wb_path.exists():
            wb_path = ROOT / "originals" / "Marksans Keshav.xlsx"
        summary = evaluate_reconciliation_summary(wb_path)
        metrics = {m["metric"]: m for m in summary}
        self.assertEqual(metrics["Market Return (Rm)"]["status"], "verified_identity")
        self.assertEqual(metrics["Cost of Equity (Ke)"]["status"], "verified_identity")
        self.assertEqual(metrics["Post-tax Cost of Debt (Kd)"]["status"], "verified_identity")
        self.assertEqual(metrics["WACC (Target-Weighted)"]["status"], "verified_identity")
        self.assertEqual(metrics["Target Weight of Equity"]["status"], "verified_identity")
        self.assertEqual(metrics["Target Weight of Debt"]["status"], "verified_identity")
        self.assertEqual(metrics["FY25 Current Ratio"]["status"], "calculated_from_workbook")
        self.assertEqual(metrics["FY25 Quick Ratio"]["status"], "calculated_from_workbook")
        self.assertEqual(metrics["Risk-Free Rate (Rf)"]["status"], "sourced_assumption")
        self.assertEqual(metrics["Pre-tax Cost of Debt"]["status"], "sourced_assumption")
        self.assertEqual(metrics["Effective Tax Rate"]["status"], "analyst_assumption")

    def test_non_circular_rm_verification_and_strict_numeric_parsing(self):
        """Regression test: verify non-circular Rm evaluation and strict rejection of None/non-numeric."""
        from build_package import evaluate_reconciliation_summary, _parse_finite_float
        wb_path = ROOT / "candidates" / "run_native_baseline_v5" / "Marksans-Recalculated.xlsx"
        if not wb_path.exists():
            wb_path = ROOT / "candidates" / "Marksans-Recalculated.xlsx"
        if not wb_path.exists():
            wb_path = ROOT / "originals" / "Marksans Keshav.xlsx"

        summary = evaluate_reconciliation_summary(wb_path)
        metrics = {m["metric"]: m for m in summary}
        # Non-circular check: exact_value should match RM!F12 (0.163692)
        self.assertEqual(metrics["Market Return (Rm)"]["exact_value"], "0.163692")
        self.assertEqual(metrics["Market Return (Rm)"]["status"], "verified_identity")

        # Test strict parsing rejects None, bool, non-numeric, or NaN/Inf
        with self.assertRaises(ValueError):
            _parse_finite_float(None, "Test!A1")
        with self.assertRaises(ValueError):
            _parse_finite_float(True, "Test!A1")
        with self.assertRaises(ValueError):
            _parse_finite_float("not_a_number", "Test!A1")
        with self.assertRaises(ValueError):
            _parse_finite_float(float("nan"), "Test!A1")
        with self.assertRaises(ValueError):
            _parse_finite_float(float("inf"), "Test!A1")

    def test_second_company_cipla_request_validation_and_provenance(self):
        """Verify second company (Cipla) request adheres to schema, separates facts, and prevents manufactured inputs."""
        from build_package import validate_package_request
        cipla_req_path = ROOT / "examples" / "cipla-package.json"
        self.assertTrue(cipla_req_path.exists())
        req = json.loads(cipla_req_path.read_text(encoding="utf-8"))

        # Verify company identity and accounting conventions
        self.assertEqual(req["company"]["name"], "Cipla Limited")
        self.assertEqual(req["company"]["ticker"], "CIPLA")
        self.assertEqual(req["company"]["currency"], "INR")
        self.assertEqual(req["company"]["reporting_unit"], "INR_Cr")
        self.assertEqual(req["company"]["fiscal_year_convention"], "April-March")

        # Verify source data segregation
        src_data = req["company"]["source_data"]
        self.assertIn("user_supplied_facts", src_data)
        self.assertIn("source_extracted_facts", src_data)
        self.assertIn("analyst_assumptions", src_data)
        self.assertIn("calculated_outputs", src_data)
        self.assertIn("unresolved_inputs", src_data)

        # Confirm evidence report exists
        ev_report = ROOT / "docs" / "SECOND_COMPANY_EVIDENCE_REPORT.md"
        self.assertTrue(ev_report.exists())
        content = ev_report.read_text(encoding="utf-8")
        self.assertIn("Cipla Limited", content)
        self.assertIn("compdata!C8", content)
        self.assertIn("WACC", content)

        # Negative test: verify validate_package_request rejects retired presentation field
        bad_req = copy.deepcopy(req)
        bad_req["presentation"] = {"template": "originals/ERR-4.pptx"}
        with self.assertRaises(ValueError) as ctx:
            validate_package_request(bad_req, cipla_req_path)
        self.assertIn("The 'presentation' request field has been retired", str(ctx.exception))

    def test_invalid_reporting_unit_rejected(self):
        """Negative test: verify validate_package_request rejects unsupported units."""
        from build_package import validate_package_request
        cipla_req_path = ROOT / "examples" / "cipla-package.json"
        req = json.loads(cipla_req_path.read_text(encoding="utf-8"))

        bad_req = copy.deepcopy(req)
        bad_req["company"]["reporting_unit"] = "USD_Trillion"
        with self.assertRaises(ValueError) as ctx:
            validate_package_request(bad_req, cipla_req_path)
        self.assertIn("Unsupported reporting_unit", str(ctx.exception))

    def test_build_package_cipla_end_to_end_skip_com(self):
        """Verify Cipla package builds successfully end-to-end with valid artifacts and manifest."""
        out_dir = self.test_dir / "output_cipla"
        req_path = ROOT / "examples" / "cipla-package.json"
        res = build_package(
            request_path=req_path,
            output_dir=out_dir,
            allow_proposal=True,
            force=False,
            skip_com=True,
        )
        self.assertEqual(res["status"], "validating")
        self.assertEqual(res["package_id"], "cipla-overview-fy25")
        self.assertTrue((out_dir / "Cipla-Recalculated.xlsx").exists())
        self.assertTrue((out_dir / "package.manifest.json").exists())

        manifest = json.loads((out_dir / "package.manifest.json").read_text(encoding="utf-8"))
        self.assertEqual(manifest["package_id"], "cipla-overview-fy25")
        self.assertEqual(manifest["workbook_validation"]["status"], "recalculated_checks_passed")

    def test_template_v16_reproducible_batch_build_skip_com(self):
        """Verify Template V16 reproducible batch builder reproduces all 4 targets from committed inputs."""
        import openpyxl
        from work.build_template_v16_models import build_all_reproducible_models, TARGETS, EXPECTED_MASTER_SHA

        repro_dir = self.test_dir / "repro_v16_test"
        summary = build_all_reproducible_models(output_dir=repro_dir, skip_com=True)

        self.assertEqual(summary["master_template_sha256"], EXPECTED_MASTER_SHA)
        self.assertTrue((repro_dir / "reproducibility_summary.json").exists())

        for target in TARGETS:
            wb_path = repro_dir / target / f"{target}-FullTemplate.xlsx"
            self.assertTrue(wb_path.exists(), f"Missing reproduced workbook for {target}")
            wb = openpyxl.load_workbook(wb_path, data_only=False)
            self.assertEqual(len(wb.sheetnames), 16, f"{target} must have 16 sheets")

            # Verify balance check formulas
            bs = wb["BS"]
            for col in ["C", "D", "E", "F", "G", "H", "I", "J", "K", "L", "M", "N"]:
                self.assertEqual(bs[f"{col}40"].value, f"={col}38-{col}16")

            # Verify Cash Flow N31 formula
            cf = wb["Cash Flow"]
            self.assertEqual(cf["N31"].value, "=N5+N14+N23")

            # Verify zero comments
            for ws in wb.worksheets:
                if hasattr(ws, "_comments"):
                    self.assertEqual(len(ws._comments), 0)


if __name__ == "__main__":
    unittest.main()
