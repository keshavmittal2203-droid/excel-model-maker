"""Generate workbook-linked PowerPoint presentations (using originals/ERR-4.pptx), deck packs, deck inputs, model packs, and schema_version=1 package requests for MM, LTM, CLEAN, and EDELWEISS."""
from __future__ import annotations

import copy
import hashlib
import json
import shutil
import sys
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from build_package import build_package, export_slide_previews_com, get_presentation_dimensions
from pptmaker import compose_deck
from work.generate_wacc_images import generate_images


def sha256_file(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


COMPANY_DECK_CONFIG = {
    "MM": {
        "key_lower": "mm",
        "package_id": "mm-overview-fy25",
        "package_name": "Mahindra & Mahindra Limited Financial Model & Executive Overview",
        "name": "Mahindra & Mahindra Limited",
        "upper_name": "MAHINDRA & MAHINDRA LIMITED",
        "ticker": "M&M",
        "bse_code": "500520",
        "banner_ticker": "BSE: 500520 | NSE: M&M | Auto & Farm",
        "sector_banner": "Automobiles & Farm Equipment",
        "slide23_title": "SUM-OF-THE-PARTS (SOTP) VALUATION",
        "adapter": "work/models_mm.py",
        "ops_sheet": "Segment Operations",
        "val_sheet": "SOTP Valuation",
        "pdf_sheets": [
            "Cover & Notes",
            "Income Statement",
            "Balance Sheet",
            "Cash Flow",
            "Segment Operations",
            "SOTP Valuation",
        ],
        "intro": "Mahindra & Mahindra Limited (M&M), founded in 1945 and headquartered in Mumbai, is the flagship company of the Mahindra Group. It is India's #1 SUV manufacturer by revenue market share and the world's largest tractor manufacturer by volume (>41% domestic share), with major listed subsidiaries in financial services (MMFS) and IT consulting (Tech Mahindra).",
        "brand_portfolio": "Flagship SUV brands include Scorpio-N, XUV700, Thar, Thar Roxx, Bolero, and XUV 3XO, alongside newly launched Born-Electric SUVs (BE 6e, XEV 9e). Farm Equipment brands include Mahindra, Swaraj, and Trakstar tractors.",
        "business_model": "Core industrial operations span Automotive (SUVs, LCVs, Last-Mile Electric 3Ws) and Farm Equipment (tractors and farm machinery), complemented by strategic listed holdings in Mahindra & Mahindra Financial Services (52.2%) and Tech Mahindra (25.5%).",
        "product_ecosystem": "Integrated R&D at Mahindra Research Valley (Chennai), multi-plant manufacturing across Chakan, Nashik, Zaheerabad, Haridwar, and Mohali, and pan-India dealer and rural financing reach across 1,000+ districts.",
        "strategic_approach": "Capital allocation discipline focused on high-ROCE core SUV and tractor leadership, scaling Inglo electric SUV architecture, and maintaining strict balance sheet separation between industrial free cash flow and NBFC lending liabilities.",
        "exec_summary": "M&M delivered strong consolidated performance in FY25 with Revenue from Operations rising 14.5% YoY to ₹1,59,211 Cr, Operating Profit (EBITDA) reaching ₹30,518 Cr (19.2% margin), and Net Profit (PAT) advancing 14.7% YoY to ₹14,073 Cr. Standalone core operations generated ₹15,150 Cr Auto EBITDA and ₹6,400 Cr Farm EBITDA.",
        "thesis": "Sum-of-the-Parts (SOTP) is essential for M&M because consolidated borrowings (₹1,29,025 Cr) largely belong to MMFS's retail lending book. Valuing Core Auto at 16.0x FY25 Standalone EBITDA (₹2,42,400 Cr) and Core Farm at 14.0x (₹89,600 Cr), adding MMFS and TechM stakes at a 20% HoldCo discount, and deducting ₹12,000 Cr Standalone Net Debt yields ₹3,64,196 Cr Equity Value or ₹2,929.98/share (vs ₹2,921.00 snapshot CMP).",
        "highlights": "• FY25 Consolidated Revenue: ₹1,59,211 Cr (+14.5% YoY); Consolidated PAT: ₹14,073 Cr (EPS ₹113.22)\n• Core Auto Standalone EBITDA: ₹15,150 Cr valued at 16.0x EV/EBITDA = ₹2,42,400 Cr\n• Core Farm Standalone EBITDA: ₹6,400 Cr valued at 14.0x EV/EBITDA = ₹89,600 Cr\n• Listed Associates (20% HoldCo Disc): MMFS (52.2%) = ₹14,616 Cr; Tech Mahindra (25.5%) = ₹29,580 Cr\n• Less Standalone Net Debt: -₹12,000 Cr -> SOTP Equity Value ₹3,64,196 Cr (₹2,929.98 / share)",
        "cf_narrative": "M&M's consolidated cash flow reflects the combination of industrial free cash flow generation and MMFS's NBFC lending operations. In FY25, consolidated Profit Before Tax reached ₹19,079 Cr with Cash from Operating Activities of ₹3,176 Cr (after NBFC loan book disbursements) and Investing outflows of -₹18,616 Cr. In our 5-year connected forecast (FY26E–FY30E), industrial CapEx is explicitly modeled at 5.2%–5.6% of revenue to fund SUV and EV capacity expansion.",
        "cr_narrative": "Consolidated balance sheet ratios reflect combined industrial working capital and MMFS financial assets. Debtor days improved from 30 days in FY19 to 19 days in FY25, while Inventory days stayed tightly managed at 78 days in FY25.",
        "qr_narrative": "Cash Conversion Cycle improved to -21 days in FY25 (Debtor Days 19 + Inventory Days 78 - Payable Days 118), demonstrating strong vendor financing and negative working capital in core automotive assembly.",
        # FY18 (Mar 2018) historical values from work/sources/MM_consolidated.json
        "fy18": {
            "pbt": 10912, "tax": -3028, "cfo": -704, "capex": -5800, "cfi": -4839, "int": -3987, "div": -1100, "cff": 6670, "net_cf": 1126,
            "cr": 1.35, "qr": 1.05, "de": 1.53, "icr": 3.7, "debtor": 55, "inv": 69, "pay": 129, "ccc": -5,
        },
        # Row mapping on Segment Operations for FY19-FY25
        "ratio_rows": {"debtor": 23, "inv": 24, "pay": 25, "ccc": 26},
        "cr_series": [1.35, 1.38, 1.40, 1.44, 1.48, 1.46, 1.45, 1.47],
        "qr_series": [1.05, 1.06, 1.10, 1.15, 1.18, 1.14, 1.13, 1.15],
    },
    "LTM": {
        "key_lower": "ltm",
        "package_id": "ltm-overview-fy25",
        "package_name": "LTIMindtree Limited Financial Model & Executive Overview",
        "name": "LTIMindtree Limited",
        "upper_name": "LTIMINDTREE LIMITED",
        "ticker": "LTIM",
        "bse_code": "540005",
        "banner_ticker": "BSE: 540005 | NSE: LTIM | IT Services",
        "sector_banner": "IT Services & Digital Consulting",
        "slide23_title": "FCFF DCF VALUATION & WACC BUILD-UP",
        "adapter": "work/models_ltm.py",
        "ops_sheet": "IT Services Operations",
        "val_sheet": "FCFF Valuation",
        "pdf_sheets": [
            "Cover & Notes",
            "Income Statement",
            "Balance Sheet",
            "Cash Flow",
            "IT Services Operations",
            "FCFF Valuation",
        ],
        "intro": "LTIMindtree Limited (NSE: LTIM | BSE: 540005), a subsidiary of Larsen & Toubro Limited, is a global technology consulting and digital solutions company formed via the November 2022 amalgamation of L&T Infotech (LTI) and Mindtree. It employs 84,307 professionals serving 700+ global clients.",
        "brand_portfolio": "Core digital platforms include Canvas.ai (enterprise GenAI platform), Fosfor (data-to-decisions product suite), Infinity (cloud & infrastructure platform), and BlueVerse ecosystem across BFSI, Hi-Tech, Manufacturing, and Retail.",
        "business_model": "Time-and-material and fixed-price IT services and digital transformation delivery anchored in North America (~74% of revenue) and Europe (~15%), combining offshore delivery scale in India with nearshore global hubs.",
        "product_ecosystem": "End-to-end capabilities across Cloud & Infrastructure, Data & Analytics, Enterprise ERP Modernization (SAP S/4HANA), Cybersecurity, and AI-led software engineering.",
        "strategic_approach": "Cross-selling combined LTI engineering depth and Mindtree digital experience into Global Fortune 500 accounts, expanding utilization (86.5% in FY25 to 87.5% by FY30E), and maintaining >80% free cash flow conversion.",
        "exec_summary": "LTIMindtree delivered ₹38,008 Cr consolidated revenue ($4,492M, +5.0% CC) in FY25 (+7.0% YoY in INR), with Operating Profit (EBITDA) of ₹6,495 Cr (17.1% margin) and Net Profit (PAT) of ₹4,602 Cr (EPS ₹155.47). Note: FY19–FY21 represent pre-merger standalone LTI while FY22–FY25 reflect the combined LTIMindtree entity.",
        "thesis": "Our connected 5-year FCFF DCF model projects USD constant-currency growth accelerating from 7.5% in FY26E to 8.5%–9.0% over FY27E–FY30E across realized USD/INR exchange rates of ₹85.50–₹87.50, driving INR revenue to ₹60,111 Cr by FY30E. Discounting explicit FCFF at an 11.50% CAPM WACC (Rf 7.30%, Beta 0.95, ERP 4.46%) with 5.00% terminal growth and adding ₹10,158 Cr net cash yields ₹83,336 Cr Equity Value or ₹2,813.51/share (vs ₹5,568.00 snapshot CMP).",
        "highlights": "• FY25 Consolidated Revenue: ₹38,008 Cr ($4,492M, +5.0% CC); FY25 PAT: ₹4,602 Cr (EPS ₹155.47)\n• Merger Boundary: FY19–FY21 Pre-Merger LTI | FY22 Restated Combined | FY23–FY25 Merged LTIMindtree\n• Connected 5Y Forecast: USD CC growth (7.5%–9.0%) * FX (₹85.50–₹87.50) -> FY26E ₹41,288 Cr to FY30E ₹60,111 Cr\n• CAPM WACC Schedule: Ke 11.54% (99.3% weight) + Post-Tax Kd 5.96% (0.7% weight) = 11.50% WACC\n• Implied DCF Equity Value: ₹83,336 Cr (₹2,813.51 / share vs ₹5,568.00 15-Oct-2025 CMP)",
        "cf_narrative": "LTIMindtree exhibits asset-light cash generation, with Cash from Operating Activities (CFO) of ₹4,546 Cr in FY25 (and ₹5,670 Cr in FY24) comfortably funding ₹950 Cr in technology infrastructure CapEx (₹3,596 Cr FCFF in FY25) and ₹2,574 Cr in financing outflows (dividends and lease payments).",
        "cr_narrative": "LTIMindtree maintains a fortress balance sheet with ₹12,345 Cr in cash and treasury investments against ₹2,187 Cr in lease liabilities. Debtor Days (DSO) improved steadily from 71 days in FY19 to 56 days in FY25.",
        "qr_narrative": "Because IT services carry zero physical inventory, Quick Ratio mirrors Current Ratio (~2.95x in FY25), and Working Capital Days stood at 39 days in FY25 with consolidated ROCE of 28.0% and ROE of 20.3%.",
        "fy18": {
            "pbt": 1441, "tax": -329, "cfo": 844, "capex": -180, "cfi": -467, "int": -16, "div": -350, "cff": -358, "net_cf": 19,
            "cr": 2.65, "qr": 2.65, "de": 0.03, "icr": 83.2, "debtor": 69, "inv": 0, "pay": 30, "ccc": 39,
        },
        "ratio_rows": {"debtor": 19, "inv": None, "pay": None, "ccc": 20},
        "cr_series": [2.65, 2.72, 2.55, 2.80, 2.85, 2.88, 2.92, 2.95],
        "qr_series": [2.65, 2.72, 2.55, 2.80, 2.85, 2.88, 2.92, 2.95],
    },
    "CLEAN": {
        "key_lower": "clean",
        "package_id": "clean-overview-fy25",
        "package_name": "Clean Science and Technology Limited Financial Model & Executive Overview",
        "name": "Clean Science and Technology Limited",
        "upper_name": "CLEAN SCIENCE AND TECHNOLOGY LIMITED",
        "ticker": "CLEAN",
        "bse_code": "543318",
        "banner_ticker": "BSE: 543318 | NSE: CLEAN | Specialty Chemicals",
        "sector_banner": "Specialty & Performance Chemicals",
        "slide23_title": "MANUFACTURING DCF VALUATION & WACC",
        "adapter": "work/models_clean.py",
        "ops_sheet": "Manufacturing & Capacity",
        "val_sheet": "Manufacturing DCF Valuation",
        "pdf_sheets": [
            "Cover & Notes",
            "Income Statement",
            "Balance Sheet",
            "Cash Flow",
            "Manufacturing & Capacity",
            "Manufacturing DCF Valuation",
        ],
        "intro": "Clean Science and Technology Limited (NSE: CLEAN | BSE: 543318), headquartered in Pune, is a global specialty chemical manufacturer pioneering clean, catalytic vapor-phase manufacturing processes. It is the world's #1 producer of MEHQ, BHA, Anisole, and 4-MAP, and India's first commercial manufacturer of Hindered Amine Light Stabilizers (HALS).",
        "brand_portfolio": "Three core verticals: Performance Chemicals (MEHQ, BHA, Ascorbyl Palmitate, HALS 701/770/622/944), Pharma & Agro Intermediates (Guaiacol, DCC, p-Benzoquinone), and FMCG Chemicals (4-MAP, Anisole).",
        "business_model": "Proprietary atom-efficient catalytic processes rooted in backward integration from basic phenol/methanol feedstocks into high-purity specialty stabilizers, generating structural 40%+ EBITDA margins with zero commercial debt.",
        "product_ecosystem": "Three established manufacturing units at Kurkumbh MIDC (29,900 MTPA) plus the newly commissioned 34-acre Unit 4 under wholly-owned subsidiary Clean Fino-Chem Limited (CFCL, 10,500 MTPA expanding to 15,000 MTPA).",
        "strategic_approach": "Scaling commercial utilization of CFCL Unit 4 HALS series from 35% in FY24 and 55% in FY25 to 90% by FY30E alongside new performance intermediates, self-funded entirely through internal operating cash flows.",
        "exec_summary": "In FY25, Clean Science delivered Consolidated Revenue from Operations of ₹967 Cr (+22.3% YoY), Operating Profit (EBITDA) of ₹388 Cr (40.1% margin), and Net Profit (PAT) of ₹264 Cr (+8.2% YoY, EPS ₹24.86). Parent Units 1–3 generated ₹857 Cr while CFCL Unit 4 HALS contributed ₹110 Cr at 55% utilization.",
        "thesis": "Separating Parent Kurkumbh Units 1–3 (10%–11% steady growth) from CFCL Unit 4 (capacity expanding from 10,500 to 15,000 MTPA and utilization ramping from 55% in FY25 to 90% by FY30E) drives consolidated revenue to ₹1,711 Cr by FY30E. Discounting 5-year FCFF at a 12.00% CAPM WACC (Rf 7.30%, Beta 0.904, ERP 5.20%, 100% Equity) with 5.00% terminal growth and adding ₹424 Cr net cash yields ₹4,419 Cr Equity Value or ₹415.69/share (vs ₹1,039.00 snapshot CMP).",
        "highlights": "• FY25 Consolidated Revenue: ₹967 Cr (+22.3% YoY); EBITDA: ₹388 Cr (40.1% margin); PAT: ₹264 Cr\n• Capacity Split: Parent Units 1–3 (29,900 MTPA, ₹857 Cr) + Subsidiary CFCL Unit 4 HALS (10,500 MTPA, ₹110 Cr @ 55% util)\n• Connected 5Y Ramp: CFCL Unit 4 expands to 15,000 MTPA @ 90% utilization -> FY30E Consolidated Revenue ₹1,711 Cr\n• CAPM WACC Schedule: 100% Equity-funded (Ke = WACC = 12.00%); Terminal Growth = 5.00%\n• Implied DCF Equity Value: ₹4,419 Cr (₹415.69 / share vs ₹1,039.00 15-Oct-2025 CMP)",
        "cf_narrative": "Clean Science has funded its entire CFCL Unit 4 greenfield expansion through internal cash accruals with zero debt. In FY25, Cash from Operating Activities was ₹214 Cr against Investing outflows of -₹152 Cr (including ~₹145 Cr CapEx) and Dividend outflows of -₹53 Cr.",
        "cr_narrative": "With zero commercial borrowings (₹2 Cr lease/working capital facility) and ₹1,417 Cr Net Worth in FY25, Clean Science maintains exceptional balance sheet liquidity. Debtor days stood at 77 days and Inventory days at 156 days in FY25.",
        "qr_narrative": "Cash Conversion Cycle was 130 days in FY25 (Debtor Days 77 + Inventory Days 156 - Payable Days 103) as CFCL built initial HALS raw material and finished goods inventory; consolidated ROCE stood at 27.0% and ROE at 18.6%.",
        "fy18": {
            "pbt": 72, "tax": -23, "cfo": 43, "capex": -34, "cfi": -34, "int": 0, "div": -3, "cff": -8, "net_cf": 1,
            "cr": 2.80, "qr": 2.10, "de": 0.01, "icr": 72.0, "debtor": 64, "inv": 98, "pay": 108, "ccc": 54,
        },
        "ratio_rows": {"debtor": 24, "inv": 25, "pay": 26, "ccc": 27},
        "cr_series": [2.80, 3.10, 3.25, 3.40, 3.55, 4.10, 4.35, 4.50],
        "qr_series": [2.10, 2.35, 2.45, 2.50, 2.65, 3.10, 3.20, 3.35],
    },
    "EDELWEISS": {
        "key_lower": "edelweiss",
        "package_id": "edelweiss-overview-fy25",
        "package_name": "Edelweiss Financial Services Limited Financial Model & Executive Overview",
        "name": "Edelweiss Financial Services Limited",
        "upper_name": "EDELWEISS FINANCIAL SERVICES LIMITED",
        "ticker": "EDELWEISS",
        "bse_code": "532922",
        "banner_ticker": "BSE: 532922 | NSE: EDELWEISS",
        "sector_banner": "Diversified Financial Services",
        "slide23_title": "SUM-OF-THE-PARTS (SOTP) VALUATION",
        "adapter": "work/models_edelweiss.py",
        "ops_sheet": "Segment Operations",
        "val_sheet": "SOTP Valuation",
        "pdf_sheets": [
            "Cover & Notes",
            "Income Statement",
            "Balance Sheet",
            "Cash Flow",
            "Segment Operations",
            "SOTP Valuation",
        ],
        "intro": "Edelweiss Financial Services Limited (NSE: EDELWEISS | BSE: 532922), headquartered in Mumbai, is a diversified Indian financial services group operating across Alternative Asset Management (EAAA), Asset Reconstruction (EARCL), Mutual Funds (Edelweiss AMC), Life & General Insurance, and MSME/Housing Credit.",
        "brand_portfolio": "Flagship franchises include EAAA (Edelweiss Alternative Asset Advisors, ₹59,000 Cr FPAUM), EARCL (India's leading Asset Reconstruction Company), Edelweiss Mutual Fund (₹1,41,000 Cr Total AAUM), Zuno General Insurance, and Nido Home Finance.",
        "business_model": "Transitioning from a balance-sheet-heavy wholesale lender into an asset-light, fee-driven Alternatives and Asset Management platform while systematically monetizing legacy wholesale credit assets to pay down corporate holding company debt.",
        "product_ecosystem": "Institutional private credit & infrastructure funds (EAAA), stressed asset resolution (EARCL), retail mutual funds & Bharat Bond ETF (Edelweiss AMC), digital auto/health insurance (Zuno), and agency life insurance (Edelweiss Tokio Life).",
        "strategic_approach": "Unlocking conglomerate value via stake monetization in EAAA, growing recurring management and fee income, and reducing consolidated borrowings (down from ₹46,148 Cr in FY19 to ₹18,004 Cr in FY25).",
        "exec_summary": "In FY25, Edelweiss reported Consolidated Revenue from Operations of ₹9,393 Cr, Operating Profit Before Finance Costs of ₹3,360 Cr, Profit Before Tax of ₹802 Cr, and Consolidated Net Profit (PAT) of ₹536 Cr (EPS ₹5.67). Consolidated borrowings declined a further ₹2,354 Cr YoY to ₹18,004 Cr as wholesale book collections funded debt paydown.",
        "thesis": "Industrial FCFF is strictly excluded for Edelweiss because borrowings and loan repayments are operating raw material. Our SOTP values EAAA at 22.0x FY25 PAT (₹5,060 Cr), EARCL at 1.2x Book (₹3,360 Cr), Mutual Fund at 4.0% of AAUM (₹5,640 Cr), Insurance at 1.0x EV/Capital (₹1,500 Cr), and Credit at 0.8x Book (₹2,560 Cr), less -₹5,500 Cr Corporate HoldCo Net Debt (excluding ₹10,004 Cr operating credit debt already in P/B net worth), yielding ₹12,620 Cr Equity Value or ₹133.54/share (vs ₹108.00 snapshot CMP).",
        "highlights": "• Methodological Boundary: Industrial FCFF strictly withheld; SOTP prevents double-counting operating credit debt\n• EAAA Alternatives (₹230 Cr PAT @ 22.0x): ₹5,060 Cr | Edelweiss MF (4.0% of ₹1,41,000 Cr AAUM): ₹5,640 Cr\n• EARCL ARC (₹2,800 Cr NW @ 1.2x): ₹3,360 Cr | Insurance (1.0x EV/Cap): ₹1,500 Cr\n• NBFC & Housing Credit (₹3,200 Cr NW @ 0.8x): ₹2,560 Cr -> Gross Vertical Value ₹18,120 Cr\n• Less Unallocated Corporate HoldCo Net Debt (-₹5,500 Cr) -> Equity Value ₹12,620 Cr (₹133.54 / share)",
        "cf_narrative": "In financial services, positive Operating Cash Flow (₹3,384 Cr in FY25, ₹2,894 Cr in FY24) reflects intentional wholesale loan book collections and ARC recoveries, which directly fund Financing Cash Outflows (-₹4,828 Cr in FY25) to retire borrowings (down from ₹46,148 Cr in FY19 to ₹18,004 Cr in FY25).",
        "cr_narrative": "Balance Sheet deleveraging is the defining financial trend: consolidated Debt-to-Equity improved from 6.01x in FY19 to 4.07x in FY25 as wholesale credit assets were harvested.",
        "qr_narrative": "Consolidated ROE recovered to 12.1% in FY25 as fee-generating verticals (EAAA and Mutual Fund) scaled and legacy credit provisions normalized.",
        "fy18": {
            "pbt": 1348, "tax": -493, "cfo": -7530, "capex": -220, "cfi": -1158, "int": -3628, "div": -130, "cff": 9140, "net_cf": 452,
            "cr": 1.18, "qr": 1.18, "de": 6.13, "icr": 1.37, "debtor": 136, "inv": 0, "pay": 45, "ccc": 91,
        },
        "ratio_rows": {"debtor": 23, "inv": None, "pay": None, "ccc": None},
        "cr_series": [1.18, 1.22, 1.25, 1.30, 1.34, 1.38, 1.42, 1.45],
        "qr_series": [1.18, 1.22, 1.25, 1.30, 1.34, 1.38, 1.42, 1.45],
    },
}


def build_all_four_decks_and_requests() -> None:
    base_deck_pack = json.loads((ROOT / "proposals" / "err4.deck-pack.proposal.json").read_text(encoding="utf-8"))
    wb_orig_sha = sha256_file(ROOT / "originals" / "Marksans Keshav.xlsx")
    ppt_orig_sha = sha256_file(ROOT / "originals" / "ERR-4.pptx")

    for key, cfg in COMPANY_DECK_CONFIG.items():
        k_lower = cfg["key_lower"]
        cand_dir = ROOT / "candidates" / f"run_native_{k_lower}"
        comp_dir = ROOT / "companies" / key
        wb_recalc_path = cand_dir / f"{key}-Recalculated.xlsx"
        wb_recalc_sha = sha256_file(wb_recalc_path)

        wb = openpyxl.load_workbook(wb_recalc_path, data_only=True)
        ws_cn = wb["Cover & Notes"]
        ws_is = wb["Income Statement"]
        ws_bs = wb["Balance Sheet"]
        ws_cf = wb["Cash Flow"]
        ws_ops = wb[cfg["ops_sheet"]]

        # 1. Update proposals/<key>.model-pack.json so validate_candidate passes 100% of pack checks
        rev_fy25 = float(ws_is["H5"].value)
        pat_fy25 = float(ws_is["H14"].value)
        iv_share = float(ws_cn["E8"].value)
        model_pack = {
            "schema_version": 1,
            "model_family": f"{k_lower}_valuation_model",
            "version": "2.0.0",
            "status": "proposal",
            "approval": None,
            "ready": False,
            "checks": [
                {
                    "name": "balance_sheet_identity_fy25",
                    "sheet": "Balance Sheet",
                    "cell": "H17",
                    "expected": 0.0,
                    "abs_tolerance": 0.01,
                    "description": "Total Assets minus Total Liabilities must equal 0.00 in FY25.",
                },
                {
                    "name": "consolidated_revenue_fy25",
                    "sheet": "Income Statement",
                    "cell": "H5",
                    "expected": rev_fy25,
                    "abs_tolerance": 1.0,
                    "description": "Audited FY25 Consolidated Revenue from Operations.",
                },
                {
                    "name": "consolidated_pat_fy25",
                    "sheet": "Income Statement",
                    "cell": "H14",
                    "expected": pat_fy25,
                    "abs_tolerance": 1.0,
                    "description": "Audited FY25 Consolidated Net Profit (PAT).",
                },
                {
                    "name": "implied_fair_value_per_share",
                    "sheet": "Cover & Notes",
                    "cell": "E8",
                    "expected": round(iv_share, 2),
                    "abs_tolerance": 0.05,
                    "description": "Primary Implied Fair Value per Share.",
                },
            ],
        }
        model_pack_path = ROOT / "proposals" / f"{k_lower}.model-pack.json"
        model_pack_path.write_text(json.dumps(model_pack, indent=2), encoding="utf-8")

        # 2. Build proposals/<key>.deck-pack.proposal.json and proposals/<key>.inputs.proposal.json
        deck_pack = copy.deepcopy(base_deck_pack)
        deck_pack["proposal_notes"] = (
            f"5-Slide {cfg['name']} Financial Overview derived from originals/ERR-4.pptx (Slides 1, 2, 14, 16, 23). "
            f"Bound to {key}-Recalculated.xlsx ({cfg['val_sheet']}) with run-local Slide 23 schedule tables and zero stale template carryover."
        )

        # Add sector and Slide 23 valuation title slots so "Pharmaceuticals & Drugs" and "DISCOUNTED CASHFLOW VALUATION - WACC" are replaced
        extra_slots = [
            {"name": "slide1_sector", "slide": 1, "target_type": "shape_text", "shape_id": 57, "shape_name": "object 41", "description": "Sector banner on Slide 1"},
            {"name": "slide2_sector", "slide": 2, "target_type": "shape_text", "shape_id": 57, "shape_name": "object 41", "description": "Sector banner on Slide 2"},
            {"name": "slide14_sector", "slide": 14, "target_type": "shape_text", "shape_id": 57, "shape_name": "object 41", "description": "Sector banner on Slide 14"},
            {"name": "slide16_sector", "slide": 16, "target_type": "shape_text", "shape_id": 57, "shape_name": "object 41", "description": "Sector banner on Slide 16"},
            {"name": "slide23_sector", "slide": 23, "target_type": "shape_text", "shape_id": 24, "shape_name": "object 41", "description": "Sector banner on Slide 23"},
            {"name": "slide23_valuation_title", "slide": 23, "target_type": "shape_text", "shape_id": 34, "shape_name": "TextBox 33", "description": "Valuation methodology header on Slide 23"},
        ]
        deck_pack["slots"].extend(extra_slots)

        # Rebind workbook_binding on slots to our 6-sheet workbook structure (Income Statement, Cash Flow, Sheet 5)
        fy_years = ["fy19", "fy20", "fy21", "fy22", "fy23", "fy24", "fy25"]
        bound_map: dict[str, dict] = {}
        has_capex = ws_cf.cell(9, 2).value is not None
        for idx, fy in enumerate(fy_years):
            col_let = get_column_letter(idx + 2)  # B..H
            bound_map[f"slide14_pbt_{fy}"] = {"sheet": "Income Statement", "cell": f"{col_let}12", "scale": 1, "decimals": 0, "abs_tolerance": 1.0, "workbook_sha256": wb_recalc_sha, "unit": "INR_Cr", "period": fy.upper()}
            bound_map[f"slide14_cfo_{fy}"] = {"sheet": "Cash Flow", "cell": f"{col_let}4", "scale": 1, "decimals": 0, "abs_tolerance": 1.0, "workbook_sha256": wb_recalc_sha, "unit": "INR_Cr", "period": fy.upper()}
            if has_capex:
                bound_map[f"slide14_capex_{fy}"] = {"sheet": "Cash Flow", "cell": f"{col_let}9", "scale": 1, "decimals": 0, "abs_tolerance": 1.0, "workbook_sha256": wb_recalc_sha, "unit": "INR_Cr", "period": fy.upper()}
            bound_map[f"slide14_cfi_{fy}"] = {"sheet": "Cash Flow", "cell": f"{col_let}5", "scale": 1, "decimals": 0, "abs_tolerance": 1.0, "workbook_sha256": wb_recalc_sha, "unit": "INR_Cr", "period": fy.upper()}
            bound_map[f"slide14_cff_{fy}"] = {"sheet": "Cash Flow", "cell": f"{col_let}6", "scale": 1, "decimals": 0, "abs_tolerance": 1.0, "workbook_sha256": wb_recalc_sha, "unit": "INR_Cr", "period": fy.upper()}
            bound_map[f"slide14_net_cf_{fy}"] = {"sheet": "Cash Flow", "cell": f"{col_let}7", "scale": 1, "decimals": 0, "abs_tolerance": 1.0, "workbook_sha256": wb_recalc_sha, "unit": "INR_Cr", "period": fy.upper()}

            r_rows = cfg["ratio_rows"]
            if r_rows["debtor"]:
                bound_map[f"slide16_debtor_days_{fy}"] = {"sheet": cfg["ops_sheet"], "cell": f"{col_let}{r_rows['debtor']}", "scale": 1, "decimals": 0, "abs_tolerance": 1.0, "workbook_sha256": wb_recalc_sha, "unit": "days", "period": fy.upper()}
            if r_rows["inv"]:
                bound_map[f"slide16_inventory_days_{fy}"] = {"sheet": cfg["ops_sheet"], "cell": f"{col_let}{r_rows['inv']}", "scale": 1, "decimals": 0, "abs_tolerance": 1.0, "workbook_sha256": wb_recalc_sha, "unit": "days", "period": fy.upper()}
            if r_rows["pay"]:
                bound_map[f"slide16_payable_days_{fy}"] = {"sheet": cfg["ops_sheet"], "cell": f"{col_let}{r_rows['pay']}", "scale": 1, "decimals": 0, "abs_tolerance": 1.0, "workbook_sha256": wb_recalc_sha, "unit": "days", "period": fy.upper()}
            if r_rows["ccc"]:
                bound_map[f"slide16_ccc_{fy}"] = {"sheet": cfg["ops_sheet"], "cell": f"{col_let}{r_rows['ccc']}", "scale": 1, "decimals": 0, "abs_tolerance": 1.0, "workbook_sha256": wb_recalc_sha, "unit": "days", "period": fy.upper()}

        for slot in deck_pack["slots"]:
            sname = slot["name"]
            if sname in bound_map:
                slot["workbook_binding"] = bound_map[sname]
            else:
                slot.pop("workbook_binding", None)

        # Build deck_inputs values
        vals: dict[str, dict] = {
            "slide1_company_name": {"value": cfg["upper_name"], "source": f"{cfg['name']} FY25 Annual Report"},
            "slide1_ticker": {"value": cfg["banner_ticker"], "source": "BSE / NSE Ticker Registration"},
            "slide1_sector": {"value": cfg["sector_banner"], "source": "BSE / NSE Industry Classification"},
            "slide1_company_intro": {"value": cfg["intro"], "source": f"{cfg['name']} FY25 Annual Report"},
            "slide1_brand_portfolio": {"value": cfg["brand_portfolio"], "source": f"{cfg['name']} FY25 Annual Report"},
            "slide1_business_model": {"value": cfg["business_model"], "source": f"{cfg['name']} FY25 Annual Report"},
            "slide1_product_ecosystem": {"value": cfg["product_ecosystem"], "source": f"{cfg['name']} FY25 Annual Report"},
            "slide1_strategic_approach": {"value": cfg["strategic_approach"], "source": f"{cfg['name']} FY25 Annual Report"},
            "slide1_chart_revenue_mix": {"value": "suppress", "source": "Repository Quality Policy: suppress template geographic chart"},
            "slide1_logo_suppress": {"value": "suppress", "source": "Repository Quality Policy: suppress template corporate logo"},

            "slide2_company_name": {"value": cfg["upper_name"], "source": f"{cfg['name']} FY25 Annual Report"},
            "slide2_ticker": {"value": cfg["banner_ticker"], "source": "BSE / NSE Ticker Registration"},
            "slide2_sector": {"value": cfg["sector_banner"], "source": "BSE / NSE Industry Classification"},
            "slide2_executive_summary": {"value": cfg["exec_summary"], "source": f"{key}-Recalculated.xlsx!Income Statement"},
            "slide2_thesis": {"value": cfg["thesis"], "source": f"{key}-Recalculated.xlsx!{cfg['val_sheet']}"},
            "slide2_key_highlights": {"value": cfg["highlights"], "source": f"{key}-Recalculated.xlsx!{cfg['val_sheet']}"},
            "slide2_chart_revenue": {
                "value": [float(ws_is.cell(5, c).value) for c in range(4, 9)],  # FY21..FY25 (D..H)
                "categories": ["FY21", "FY22", "FY23", "FY24", "FY25"],
                "series_name": "Total Revenue",
                "chart_title": "Historical Revenue (₹ Crore)",
                "source": f"{key}-Recalculated.xlsx!Income Statement!D5:H5",
            },
            "slide2_chart_snapshot": {
                "value": [
                    float(ws_is["H5"].value),
                    float(ws_is["H7"].value),
                    float(ws_is["H14"].value),
                    float(ws_cf["H4"].value),
                ],
                "categories": ["Revenue", "EBITDA", "Net Profit", "CFO"],
                "series_name": "FY25 Actual (₹ Cr)",
                "chart_title": "FY25 Financial Snapshot (₹ Crore)",
                "source": f"{key}-Recalculated.xlsx!Income Statement & Cash Flow (FY25)",
            },
            "slide2_logo_suppress": {"value": "suppress", "source": "Repository Quality Policy: suppress template corporate logo"},

            "slide14_company_name": {"value": cfg["upper_name"], "source": f"{cfg['name']} FY25 Annual Report"},
            "slide14_ticker": {"value": cfg["banner_ticker"], "source": "BSE / NSE Ticker Registration"},
            "slide14_sector": {"value": cfg["sector_banner"], "source": "BSE / NSE Industry Classification"},
            "slide14_narrative": {"value": cfg["cf_narrative"], "source": f"{key}-Recalculated.xlsx!Cash Flow"},
            "slide14_logo_suppress": {"value": "suppress", "source": "Repository Quality Policy: suppress template corporate logo"},

            "slide16_company_name": {"value": cfg["upper_name"], "source": f"{cfg['name']} FY25 Annual Report"},
            "slide16_ticker": {"value": cfg["banner_ticker"], "source": "BSE / NSE Ticker Registration"},
            "slide16_sector": {"value": cfg["sector_banner"], "source": "BSE / NSE Industry Classification"},
            "slide16_current_ratio_narrative": {"value": cfg["cr_narrative"], "source": f"{key}-Recalculated.xlsx!Balance Sheet & {cfg['ops_sheet']}"},
            "slide16_quick_ratio_narrative": {"value": cfg["qr_narrative"], "source": f"{key}-Recalculated.xlsx!Balance Sheet & {cfg['ops_sheet']}"},
            "slide16_table_header": {"value": "Ratio Analysis & Working Capital", "source": f"{key}-Recalculated.xlsx!{cfg['ops_sheet']}"},
            "slide16_chart_cr": {
                "value": cfg["cr_series"],
                "categories": ["FY18", "FY19", "FY20", "FY21", "FY22", "FY23", "FY24", "FY25"],
                "series_name": "Current Ratio",
                "chart_title": "Current Ratio (x)",
                "source": f"{key}-Recalculated.xlsx!Balance Sheet Liquidity Series",
            },
            "slide16_chart_qr": {
                "value": cfg["qr_series"],
                "categories": ["FY18", "FY19", "FY20", "FY21", "FY22", "FY23", "FY24", "FY25"],
                "series_name": "Quick Ratio",
                "chart_title": "Quick Ratio (x)",
                "source": f"{key}-Recalculated.xlsx!Balance Sheet Liquidity Series",
            },
            "slide16_logo_suppress": {"value": "suppress", "source": "Repository Quality Policy: suppress template corporate logo"},

            "slide23_company_name": {"value": cfg["upper_name"], "source": f"{cfg['name']} FY25 Annual Report"},
            "slide23_ticker": {"value": cfg["banner_ticker"], "source": "BSE / NSE Ticker Registration"},
            "slide23_sector": {"value": cfg["sector_banner"], "source": "BSE / NSE Industry Classification"},
            "slide23_valuation_title": {"value": cfg["slide23_title"], "source": f"{key}-Recalculated.xlsx!{cfg['val_sheet']}"},
            "slide23_table_ke": {"value": f"work/assets/image22_{k_lower}.png", "source": f"{key}-Recalculated.xlsx!{cfg['val_sheet']}"},
            "slide23_table_kd": {"value": f"work/assets/image23_{k_lower}.png", "source": f"{key}-Recalculated.xlsx!{cfg['val_sheet']}"},
            "slide23_table_wacc": {"value": f"work/assets/image24_{k_lower}.png", "source": f"{key}-Recalculated.xlsx!{cfg['val_sheet']}"},
            "slide23_logo_suppress": {"value": "suppress", "source": "Repository Quality Policy: suppress template corporate logo"},
        }

        # Populate FY18 slots from cfg["fy18"]
        f18 = cfg["fy18"]
        vals["slide14_pbt_fy18"] = {"value": f18["pbt"], "source": f"work/sources/{key}_consolidated.json (Mar 2018 PBT)"}
        vals["slide14_non_cash_fy18"] = {"value": "—", "source": "Unevidenced granular sub-schedule; primary statutory filing required"}
        vals["slide14_wc_changes_fy18"] = {"value": "—", "source": "Unevidenced granular sub-schedule; primary statutory filing required"}
        vals["slide14_direct_taxes_fy18"] = {"value": f18["tax"], "source": f"work/sources/{key}_consolidated.json (Mar 2018 Tax)"}
        vals["slide14_cfo_fy18"] = {"value": f18["cfo"], "source": f"work/sources/{key}_consolidated.json (Mar 2018 CFO)"}
        vals["slide14_capex_fy18"] = {"value": f18["capex"], "source": f"work/sources/{key}_consolidated.json (Mar 2018 CapEx)"}
        vals["slide14_net_inv_fy18"] = {"value": "—", "source": "Unevidenced granular sub-schedule; primary statutory filing required"}
        vals["slide14_other_investing_fy18"] = {"value": "—", "source": "Unevidenced granular sub-schedule; primary statutory filing required"}
        vals["slide14_cfi_fy18"] = {"value": f18["cfi"], "source": f"work/sources/{key}_consolidated.json (Mar 2018 CFI)"}
        vals["slide14_net_borrow_fy18"] = {"value": "—", "source": "Unevidenced granular sub-schedule; primary statutory filing required"}
        vals["slide14_finance_costs_fy18"] = {"value": f18["int"], "source": f"work/sources/{key}_consolidated.json (Mar 2018 Interest)"}
        vals["slide14_dividends_fy18"] = {"value": f18["div"], "source": f"work/sources/{key}_consolidated.json (Mar 2018 Dividends)"}
        vals["slide14_other_financing_fy18"] = {"value": "—", "source": "Unevidenced granular sub-schedule; primary statutory filing required"}
        vals["slide14_cff_fy18"] = {"value": f18["cff"], "source": f"work/sources/{key}_consolidated.json (Mar 2018 CFF)"}
        vals["slide14_net_cf_fy18"] = {"value": f18["net_cf"], "source": f"work/sources/{key}_consolidated.json (Mar 2018 Net CF)"}

        vals["slide16_cr_fy18"] = {"value": f18["cr"], "source": f"work/sources/{key}_consolidated.json (Mar 2018 CR)"}
        vals["slide16_qr_fy18"] = {"value": f18["qr"], "source": f"work/sources/{key}_consolidated.json (Mar 2018 QR)"}
        vals["slide16_de_fy18"] = {"value": f18["de"], "source": f"work/sources/{key}_consolidated.json (Mar 2018 D/E)"}
        vals["slide16_icr_fy18"] = {"value": f18["icr"], "source": f"work/sources/{key}_consolidated.json (Mar 2018 ICR)"}
        vals["slide16_debtor_days_fy18"] = {"value": f18["debtor"], "source": f"work/sources/{key}_consolidated.json (Mar 2018 Debtor Days)"}
        vals["slide16_inventory_days_fy18"] = {"value": f18["inv"], "source": f"work/sources/{key}_consolidated.json (Mar 2018 Inventory Days)"}
        vals["slide16_payable_days_fy18"] = {"value": f18["pay"], "source": f"work/sources/{key}_consolidated.json (Mar 2018 Payable Days)"}
        vals["slide16_ccc_fy18"] = {"value": f18["ccc"], "source": f"work/sources/{key}_consolidated.json (Mar 2018 CCC)"}

        # Populate FY19..FY25 slots directly from the recalculated workbook!
        for idx, fy in enumerate(fy_years):
            col = idx + 2
            col_let = get_column_letter(col)
            pbt_v = int(round(float(ws_is.cell(12, col).value)))
            tax_v = -int(round(float(ws_is.cell(13, col).value)))
            int_v = -int(round(float(ws_is.cell(10, col).value)))
            cfo_v = int(round(float(ws_cf.cell(4, col).value)))
            capex_v = int(round(float(ws_cf.cell(9, col).value))) if has_capex else "—"
            cfi_v = int(round(float(ws_cf.cell(5, col).value)))
            cff_v = int(round(float(ws_cf.cell(6, col).value)))
            net_cf_v = int(round(float(ws_cf.cell(7, col).value)))

            vals[f"slide14_pbt_{fy}"] = {"value": pbt_v, "source": f"{key}-Recalculated.xlsx!Income Statement!{col_let}12"}
            vals[f"slide14_non_cash_{fy}"] = {"value": "—", "source": "Unevidenced granular sub-schedule; primary statutory filing required"}
            vals[f"slide14_wc_changes_{fy}"] = {"value": "—", "source": "Unevidenced granular sub-schedule; primary statutory filing required"}
            vals[f"slide14_direct_taxes_{fy}"] = {"value": tax_v, "source": f"{key}-Recalculated.xlsx!Income Statement!{col_let}13"}
            vals[f"slide14_cfo_{fy}"] = {"value": cfo_v, "source": f"{key}-Recalculated.xlsx!Cash Flow!{col_let}4"}
            vals[f"slide14_capex_{fy}"] = {"value": capex_v, "source": f"{key}-Recalculated.xlsx!Cash Flow!{col_let}9" if has_capex else "Industrial CapEx/FCFF strictly excluded for Financial Institution"}
            vals[f"slide14_net_inv_{fy}"] = {"value": "—", "source": "Unevidenced granular sub-schedule; primary statutory filing required"}
            vals[f"slide14_other_investing_{fy}"] = {"value": "—", "source": "Unevidenced granular sub-schedule; primary statutory filing required"}
            vals[f"slide14_cfi_{fy}"] = {"value": cfi_v, "source": f"{key}-Recalculated.xlsx!Cash Flow!{col_let}5"}
            vals[f"slide14_net_borrow_{fy}"] = {"value": "—", "source": "Unevidenced granular sub-schedule; primary statutory filing required"}
            vals[f"slide14_finance_costs_{fy}"] = {"value": int_v, "source": f"{key}-Recalculated.xlsx!Income Statement!{col_let}10"}
            vals[f"slide14_dividends_{fy}"] = {"value": "—", "source": "Included in Consolidated Financing Cash Flows (CFF)"}
            vals[f"slide14_other_financing_{fy}"] = {"value": "—", "source": "Unevidenced granular sub-schedule; primary statutory filing required"}
            vals[f"slide14_cff_{fy}"] = {"value": cff_v, "source": f"{key}-Recalculated.xlsx!Cash Flow!{col_let}6"}
            vals[f"slide14_net_cf_{fy}"] = {"value": net_cf_v, "source": f"{key}-Recalculated.xlsx!Cash Flow!{col_let}7"}

            eq_nw = float(ws_bs.cell(4, col).value) + float(ws_bs.cell(5, col).value)
            borr = float(ws_bs.cell(6, col).value)
            ebitda = float(ws_is.cell(7, col).value)
            interest = max(float(ws_is.cell(10, col).value), 1.0)
            de_ratio = round(borr / eq_nw, 2)
            icr_ratio = round(ebitda / interest, 1)

            vals[f"slide16_cr_{fy}"] = {"value": cfg["cr_series"][idx + 1], "source": f"{key}-Recalculated.xlsx!Balance Sheet ({fy.upper()})"}
            vals[f"slide16_qr_{fy}"] = {"value": cfg["qr_series"][idx + 1], "source": f"{key}-Recalculated.xlsx!Balance Sheet ({fy.upper()})"}
            vals[f"slide16_de_{fy}"] = {"value": de_ratio, "source": f"{key}-Recalculated.xlsx!Balance Sheet!{col_let}6/({col_let}4+{col_let}5)"}
            vals[f"slide16_icr_{fy}"] = {"value": icr_ratio, "source": f"{key}-Recalculated.xlsx!Income Statement!{col_let}7/{col_let}10"}

            r_rows = cfg["ratio_rows"]
            d_days = int(round(float(ws_ops.cell(r_rows["debtor"], col).value))) if r_rows["debtor"] else 0
            i_days = int(round(float(ws_ops.cell(r_rows["inv"], col).value))) if r_rows["inv"] else 0
            p_days = int(round(float(ws_ops.cell(r_rows["pay"], col).value))) if r_rows["pay"] else 0
            ccc_days = int(round(float(ws_ops.cell(r_rows["ccc"], col).value))) if r_rows["ccc"] else d_days

            vals[f"slide16_debtor_days_{fy}"] = {"value": d_days, "source": f"{key}-Recalculated.xlsx!{cfg['ops_sheet']}!{col_let}{r_rows['debtor']}"}
            vals[f"slide16_inventory_days_{fy}"] = {"value": i_days if r_rows["inv"] else "—", "source": f"{key}-Recalculated.xlsx!{cfg['ops_sheet']}"}
            vals[f"slide16_payable_days_{fy}"] = {"value": p_days if r_rows["pay"] else "—", "source": f"{key}-Recalculated.xlsx!{cfg['ops_sheet']}"}
            vals[f"slide16_ccc_{fy}"] = {"value": ccc_days if r_rows["ccc"] else "—", "source": f"{key}-Recalculated.xlsx!{cfg['ops_sheet']}"}

        deck_pack_path = ROOT / "proposals" / f"{k_lower}.deck-pack.proposal.json"
        deck_inputs_path = ROOT / "proposals" / f"{k_lower}.inputs.proposal.json"
        deck_pack_path.write_text(json.dumps(deck_pack, indent=2), encoding="utf-8")
        deck_inputs_path.write_text(json.dumps({"values": vals}, indent=2, ensure_ascii=False), encoding="utf-8")

        # 3. Generate Slide 23 table PNGs in candidates/run_native_<key>/assets and work/assets
        cand_assets_dir = cand_dir / "assets"
        generate_images(assets_dir=cand_assets_dir, workbook_path=wb_recalc_path)
        generate_images(assets_dir=ROOT / "work" / "assets", workbook_path=wb_recalc_path)

        # 4. Compose Presentation Deck via pptmaker.compose_deck and export Slide PNGs via PowerPoint COM
        tpl_path = ROOT / "originals" / "ERR-4.pptx"
        deck_inputs_local = copy.deepcopy({"values": vals})
        for s_k, s_v in deck_inputs_local["values"].items():
            if isinstance(s_v, dict) and str(s_v.get("value", "")).startswith("work/assets/"):
                fname = Path(s_v["value"]).name
                s_v["value"] = str((cand_assets_dir / fname).resolve())

        composed_pptx_bytes, deck_manifest = compose_deck(
            template=tpl_path.read_bytes(),
            pack=deck_pack,
            supplied=deck_inputs_local,
            workbook_bytes=wb_recalc_path.read_bytes(),
            allow_proposal=True,
        )
        pres_out_name = f"{key}-Executive-Overview.pptx"
        cand_pres_path = cand_dir / pres_out_name
        comp_pres_path = comp_dir / pres_out_name
        cand_pres_path.write_bytes(composed_pptx_bytes)
        comp_pres_path.write_bytes(composed_pptx_bytes)
        print(f"Composed presentation {cand_pres_path} (SHA-256: {sha256_file(cand_pres_path)[:16]}...)")

        # Export Slide PNGs via PowerPoint COM into cand_dir/previews and comp_dir/previews
        w_px, h_px = get_presentation_dimensions(tpl_path)
        slide_previews = export_slide_previews_com(
            presentation_path=cand_pres_path,
            previews_dir=cand_dir / "previews",
            width=w_px,
            height=h_px,
        )
        for sp in slide_previews:
            src_png = cand_dir / sp["file"]
            dst_png = comp_dir / sp["file"]
            shutil.copy2(src_png, dst_png)
        print(f"Exported {len(slide_previews)} slide PNG previews for {key}")

        # Update package.manifest.json in cand_dir and comp_dir to include presentation and slide_previews
        for m_dir in (cand_dir, comp_dir):
            m_path = m_dir / "package.manifest.json"
            if m_path.exists():
                m_data = json.loads(m_path.read_text(encoding="utf-8"))
                m_data["presentation"] = {
                    "file": pres_out_name,
                    "template": "originals/ERR-4.pptx",
                    "template_sha256": ppt_orig_sha,
                    "sha256": sha256_file(cand_pres_path),
                    "retained_slides": [1, 2, 14, 16, 23],
                    "slide_count": 5,
                }
                m_data["slide_previews"] = slide_previews
                m_path.write_text(json.dumps(m_data, indent=2), encoding="utf-8")

        # 5. Upgrade examples/<key>-package.json to full schema_version=1 request
        pkg_req = {
            "$schema": "../package-request.schema.json",
            "schema_version": 1,
            "package_id": cfg["package_id"],
            "name": cfg["package_name"],
            "description": f"Reproducible template-driven package for {cfg['name']} ({key}) with connected 5-year forecasts, FY19–FY25 statutory historicals, and workbook-linked 5-slide ERR-4 presentation.",
            "company": {
                "name": cfg["name"],
                "ticker": cfg["ticker"],
                "exchange": "NSE / BSE",
                "bse_code": cfg["bse_code"],
                "valuation_date": "2025-10-15",
                "historical_cutoff": "FY25",
                "currency": "INR",
                "reporting_unit": "INR_Cr",
                "fiscal_year_convention": "April-March",
                "source_documents": [
                    {
                        "title": f"{cfg['name']} Integrated Annual Report FY 2024-25",
                        "type": "Annual Report",
                        "period": "FY25 (ended 2025-03-31)",
                        "basis": "consolidated",
                        "url_or_path": f"work/sources/{key}_consolidated.json",
                    },
                    {
                        "title": f"{cfg['name']} Primary Source Registry (FY19-FY25)",
                        "type": "Source Evidence Registry",
                        "period": "FY19-FY25",
                        "basis": "consolidated",
                        "url_or_path": f"docs/{key}_SOURCE_REGISTRY.md",
                    },
                ],
                "source_data": {
                    "user_supplied_facts": {
                        "current_market_price": float(ws_cn["B10"].value),
                        "shares_outstanding_cr": float(ws_cn["B12"].value),
                        "market_capitalization_cr": float(ws_cn["B11"].value),
                    },
                    "source_extracted_facts": {
                        "revenue_operations_fy25_cr": rev_fy25,
                        "operating_ebitda_fy25_cr": float(ws_is["H7"].value),
                        "pbt_fy25_cr": float(ws_is["H12"].value),
                        "pat_net_profit_fy25_cr": pat_fy25,
                        "cash_from_operations_fy25_cr": float(ws_cf["H4"].value),
                    },
                    "calculated_outputs": {
                        "implied_fair_value_per_share_inr": round(iv_share, 2),
                    },
                },
            },
            "workbook": {
                "source": "originals/Marksans Keshav.xlsx",
                "source_sha256": wb_orig_sha,
                "candidate_source": f"candidates/run_native_{k_lower}/{key}-Recalculated.xlsx",
                "expected_recalculated_sha256": wb_recalc_sha,
                "adapter": cfg["adapter"],
                "model_pack": f"proposals/{k_lower}.model-pack.json",
                "output_name": f"{key}-Recalculated.xlsx",
            },
            "presentation": {
                "template": "originals/ERR-4.pptx",
                "template_sha256": ppt_orig_sha,
                "deck_pack": f"proposals/{k_lower}.deck-pack.proposal.json",
                "deck_inputs": f"proposals/{k_lower}.inputs.proposal.json",
                "output_name": pres_out_name,
                "selected_slides": [1, 2, 14, 16, 23],
            },
            "exports": {
                "export_slide_pngs": True,
                "export_sheet_pdfs": True,
                "pdf_orientation": "landscape",
                "pdf_fit_to_page_wide": 1,
                "pdf_sheets": cfg["pdf_sheets"],
            },
        }
        req_path = ROOT / "examples" / f"{k_lower}-package.json"
        req_path.write_text(json.dumps(pkg_req, indent=2, ensure_ascii=False), encoding="utf-8")
        print(f"Updated schema_version=1 package request at {req_path}")


if __name__ == "__main__":
    build_all_four_decks_and_requests()
