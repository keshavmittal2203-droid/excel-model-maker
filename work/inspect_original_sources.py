import sys
import re
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('originals/Marksans Keshav.xlsx', data_only=False)

url_regex = re.compile(r'https?://[^\s"\'<>]+|www\.[^\s"\'<>]+|[a-zA-Z0-9\-\.]+\.(?:com|in|org|gov|edu|net|io)[^\s"\'<>]*')
keywords = ['screener', 'yahoo', 'investing', 'bse', 'nse', 'damodaran', 'rbi', 'bloomberg', 'moneycontrol', 'capitaline', 'statutory', 'annual report']

found_sources = []
seen = set()

for sname in wb.sheetnames:
    ws = wb[sname]
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(row=r, column=c)
            coord = f"{sname}!{cell.coordinate}"
            
            # 1. Hyperlink
            if cell.hyperlink and cell.hyperlink.target:
                k = (coord, 'hyperlink', cell.hyperlink.target)
                if k not in seen:
                    seen.add(k)
                    found_sources.append({
                        'coord': coord, 'type': 'hyperlink',
                        'target': cell.hyperlink.target,
                        'text': str(cell.value)
                    })
            
            # 2. Cell comments
            if cell.comment:
                c_text = cell.comment.text
                m = url_regex.findall(c_text)
                for u in m:
                    k = (coord, 'comment_url', u)
                    if k not in seen:
                        seen.add(k)
                        found_sources.append({
                            'coord': coord, 'type': 'comment_url',
                            'target': u, 'text': c_text
                        })
            
            # 3. Cell value string
            val_str = str(cell.value or '')
            matches = url_regex.findall(val_str)
            for m in matches:
                k = (coord, 'cell_text_url', m)
                if k not in seen:
                    seen.add(k)
                    found_sources.append({
                        'coord': coord, 'type': 'cell_text_url',
                        'target': m, 'text': val_str
                    })
            
            for kw in keywords:
                if kw in val_str.lower():
                    k = (coord, f'keyword_{kw}', kw)
                    if k not in seen:
                        seen.add(k)
                        found_sources.append({
                            'coord': coord, 'type': f'keyword_{kw}',
                            'target': kw, 'text': val_str
                        })

print(f"Total source references found in original Marksans workbook: {len(found_sources)}\n")
for src in found_sources:
    txt = src['text'].replace('\n', ' ')
    if len(txt) > 80:
        txt = txt[:77] + '...'
    print(f"{src['coord']:<25} | {src['type']:<15} | target: {src['target']:<35} | text: {txt}")
