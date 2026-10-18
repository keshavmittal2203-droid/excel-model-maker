import openpyxl
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
TARGETS = ["MM", "LTM", "CLEAN", "EDELWEISS"]

print("="*60)
print("AUDITING PRESENTATION DEFECTS ACROSS ALL 4 WORKBOOKS")
print("="*60)

for t in TARGETS:
    p = ROOT / "companies" / t / "template_v6" / f"{t}-FullTemplate.xlsx"
    wb = openpyxl.load_workbook(p, data_only=False)
    
    print(f"\n================ TARGET: {t} ================")
    
    # 1. Company Branding on Main Financials
    mf = wb["Main Financials"]
    b3_val = mf["B3"].value
    b3_font = mf["B3"].font
    print(f"  Branding at Main Financials!B3: '{b3_val}' (font: {b3_font.name}, size: {b3_font.size}, bold: {b3_font.bold}, color: {b3_font.color.rgb if b3_font.color else None})")
    
    # 2. Check Comps sheet formatting & formulas
    comps = wb["Comps"]
    print("  Comps Sheet Structure:")
    for r in range(4, 12):
        col_c = comps[f"C{r}"].value
        col_d = comps[f"D{r}"].value
        col_f = comps[f"F{r}"].value
        col_h = comps[f"H{r}"].value
        col_j = comps[f"J{r}"].value
        col_k = comps[f"K{r}"].value
        col_l = comps[f"L{r}"].value
        print(f"    Row {r:2d} | C: {col_c} | D: {col_d} | F: {col_f} | H: {col_h} | J: {col_j} | K: {col_k} | L: {col_l}")
        
    print("  Comps Valuation Summary (Rows 25-35):")
    for r in range(25, 35):
        lbl = str(comps[f"C{r}"].value or comps[f"B{r}"].value or '')
        n_val = comps[f"N{r}"].value
        p_val = comps[f"P{r}"].value
        if lbl or n_val or p_val:
            print(f"    Row {r:2d} ({lbl[:25]}): N={n_val} | P={p_val}")

    # 3. Check for blue hardcodes vs black formulas
    # In financial modeling, hardcoded numbers are typically Blue, formulas Black
    blue_cells = []
    for sname in ["Model ->", "Main Financials", "Comps", "DCF 1st way"]:
        ws = wb[sname]
        for r in range(1, min(ws.max_row+1, 50)):
            for c in range(1, min(ws.max_column+1, 15)):
                cell = ws.cell(r, c)
                f = cell.font
                v = cell.value
                if f and f.color and f.color.rgb in ["000000FF", "FF0000FF", "0000FF"]:
                    blue_cells.append((sname, cell.coordinate, v, str(v).startswith("=")))
    print(f"  Blue styled cells (first 5 of {len(blue_cells)}):")
    for sname, coord, val, is_f in blue_cells[:5]:
        print(f"    [{sname}!{coord}] value={val} (formula={is_f})")

    # 4. Check for comments / notes
    comments_found = []
    for sname in wb.sheetnames:
        ws = wb[sname]
        for r in range(1, ws.max_row+1):
            for c in range(1, ws.max_column+1):
                cell = ws.cell(r, c)
                if cell.comment:
                    comments_found.append((sname, cell.coordinate, cell.comment.text.replace("\n", " ")[:60]))
    print(f"  Total cell comments found: {len(comments_found)}")
    for sname, coord, ctext in comments_found[:5]:
        print(f"    [{sname}!{coord}]: {ctext}...")
