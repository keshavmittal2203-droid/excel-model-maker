import openpyxl

targets = ['MM', 'LTM', 'CLEAN', 'EDELWEISS']

for t in targets:
    wb = openpyxl.load_workbook(f'companies/{t}/template_v6/{t}-FullTemplate.xlsx', data_only=False)
    print(f"\n{'='*70}\nVISUAL & PRESENTATION AUDIT: {t}\n{'='*70}")
    
    # 1. Branding
    mf = wb['Main Financials']
    b3_val = mf['B3'].value
    print(f"1. Company Wordmark (Main Financials!B3): {b3_val}")
    
    # Check for Marksans in any cell across all sheets
    stale_marksans = []
    for s in wb.sheetnames:
        ws = wb[s]
        for r in range(1, 50):
            for c in range(1, 20):
                val = str(ws.cell(row=r, column=c).value or "")
                if 'marksans' in val.lower():
                    stale_marksans.append((s, r, c, val))
    print(f"2. Stale Marksans strings across all 16 sheets: {len(stale_marksans)}")
    if stale_marksans:
        print(f"   Found: {stale_marksans}")
        
    # 3. Comps sheet audit
    comps = wb['Comps']
    active_peers = [comps.cell(row=r, column=2).value for r in range(9, 14)]
    cleared_peers = [comps.cell(row=r, column=2).value for r in range(14, 16)]
    print(f"3. Comps Active Peers (Rows 9-13): {active_peers}")
    print(f"   Comps Cleared Rows (Rows 14-15): {cleared_peers}")
    print(f"   Comps Target Multiples (Row 8): EV/Rev={comps['N8'].value}, EV/EBITDA={comps['O8'].value}, P/E={comps['P8'].value}")
    
    # 4. Sheet Titles and Navigation
    print(f"4. Sheet count: {len(wb.sheetnames)} (sectordata removed: {'sectordata' not in wb.sheetnames})")
    
    # 5. Check BS balance check formatting
    bs = wb['BS']
    b40 = bs['B40'].value
    n40 = bs['N40'].value
    print(f"5. Balance Sheet Row 40 Check: Label='{b40}', Formula='{n40}'")
