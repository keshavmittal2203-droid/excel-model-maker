import openpyxl
import re

wb = openpyxl.load_workbook('originals/Marksans Keshav.xlsx', data_only=False)
ds_refs = set()
for name in wb.sheetnames:
    ws = wb[name]
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            val = str(ws.cell(r, c).value)
            if 'Data Sheet' in val:
                matches = re.findall(r"'Data Sheet'!([A-Z]+)(\d+)", val)
                for col_str, row_str in matches:
                    ds_refs.add((int(row_str), name, f"{openpyxl.utils.get_column_letter(c)}{r}", val))

print(f"Total references to Data Sheet: {len(ds_refs)}")
# Group by Data Sheet row
by_row = {}
for row_num, sheet, cell, formula in ds_refs:
    by_row.setdefault(row_num, []).append((sheet, cell, formula))

for r in sorted(by_row.keys()):
    print(f"Data Sheet Row {r:2d} (e.g. {wb['Data Sheet'].cell(r, 1).value}): referenced {len(by_row[r])} times across sheets {set(s for s, _, _ in by_row[r])}")
