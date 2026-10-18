import openpyxl
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
TARGETS = ["MM", "LTM", "CLEAN", "EDELWEISS"]

print("="*60)
print("SCANNING FOR IFERROR, NA(), AND 'N/A' FORMULAS")
print("="*60)

for t in TARGETS:
    p = ROOT / "companies" / t / "template_v6" / f"{t}-FullTemplate.xlsx"
    wb = openpyxl.load_workbook(p, data_only=False)
    
    iferror_list = []
    na_list = []
    
    for sname in wb.sheetnames:
        ws = wb[sname]
        for r in range(1, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                cell = ws.cell(r, c)
                val = cell.value
                if val and str(val).startswith("="):
                    sval = str(val).upper()
                    if "IFERROR" in sval:
                        iferror_list.append((f"{sname}!{cell.coordinate}", val))
                    if '"N/A"' in sval or "NA()" in sval:
                        na_list.append((f"{sname}!{cell.coordinate}", val))
                        
    print(f"\n{t}:")
    print(f"  IFERROR occurrences: {len(iferror_list)}")
    for coord, val in iferror_list:
        print(f"    {coord}: {val}")
        
    print(f"  N/A / NA() occurrences: {len(na_list)}")
    for coord, val in na_list:
        print(f"    {coord}: {val}")
