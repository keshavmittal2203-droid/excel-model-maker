"""Scan workbooks for formulas containing hardcoded numeric literals."""

import openpyxl
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
TARGETS = ["MM", "LTM", "CLEAN", "EDELWEISS"]
VERSION = sys.argv[1] if len(sys.argv) > 1 else "template_v5"

# Regex to find cell references (e.g. A1, $A$1, 'Sheet Name'!A1, BS!$C$10)
CELL_REF_REGEX = re.compile(r"('?[a-zA-Z0-9_\s-]+'?!)?\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?")
NUMBER_REGEX = re.compile(r"(?<![A-Z_])[-+]?(?:\d*\.\d+|\d+)")

COMMON_CONSTANTS = {"0", "1", "2", "3", "4", "5", "6", "7", "8", "9", "10", "12", "100", "365", "366"}

for tgt in TARGETS:
    wb_path = ROOT / "companies" / tgt / VERSION / f"{tgt}-FullTemplate.xlsx"
    wb = openpyxl.load_workbook(wb_path, data_only=False)
    print(f"\n==================================================")
    print(f"Scanning {tgt}: {wb_path.name}")
    print(f"==================================================")
    found_literals = []

    for sname in wb.sheetnames:
        ws = wb[sname]
        for row in ws.iter_rows():
            for cell in row:
                val = cell.value
                if val and isinstance(val, str) and val.startswith("="):
                    # Strip out all cell references from formula
                    stripped = CELL_REF_REGEX.sub("", val)
                    # Find any remaining numbers
                    nums = NUMBER_REGEX.findall(stripped)
                    suspicious = [n for n in nums if n not in COMMON_CONSTANTS and not n.startswith("-1")]
                    if suspicious:
                        found_literals.append({
                            "sheet": sname,
                            "cell": cell.coordinate,
                            "formula": val,
                            "suspicious": suspicious,
                        })

    if found_literals:
        print(f"Found {len(found_literals)} formulas with numeric literals in {tgt}:")
        for item in found_literals:
            print(f"  {item['sheet']}!{item['cell']}: {item['formula']} -> Literals: {item['suspicious']}")
    else:
        print(f"Zero suspicious numeric literals found in {tgt} formulas.")
