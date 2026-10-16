import openpyxl

wb = openpyxl.load_workbook('candidates/Marksans-Recalculated.xlsx', data_only=False)
ws_ds = wb['Data Sheet']
for r in [28, 82, 83, 84, 85]:
    lbl = ws_ds.cell(r, 2).value
    vals = [ws_ds.cell(r, c).value for c in range(3, 11)]
    print(f"{r}: {vals}")

ws_cf = wb['Cash Flow']
print("\n=== Cash Flow Formulas ===")
for r in [5, 12, 13, 14, 15, 16, 17, 18, 22, 23, 25, 26, 27, 28, 29, 30, 31]:
    lbl = ws_cf.cell(r, 2).value
    vals = [ws_cf.cell(r, c).value for c in range(6, 14)]
    print(f"{r:2d}: {lbl!s:30s} {vals}")
