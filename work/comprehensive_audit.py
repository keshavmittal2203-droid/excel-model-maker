"""Comprehensive deep audit of all 4 template_v6 workbooks."""

import win32com.client as win32
from pathlib import Path
import openpyxl
import openpyxl.utils

ROOT = Path(__file__).resolve().parent.parent
TARGETS = ["MM", "LTM", "CLEAN", "EDELWEISS"]

COM_ERRORS = {
    -2146826281: "#DIV/0!",
    -2146826273: "#VALUE!",
    -2146826265: "#REF!",
    -2146826259: "#NAME?",
    -2146826252: "#NUM!",
    -2146826246: "#N/A",
    -2146826288: "#NULL!"
}

STALE_TERMS = ["marksans", "cipla", "dr reddy", "dr. reddy", "sun pharma", "glenmark", "alembic", "laurus", "granules"]

def run_comprehensive_audit():
    print("="*60)
    print("STARTING COMPREHENSIVE DEEP AUDIT (RECHECK THRICE)")
    print("="*60)

    # 1. OpenPyXL inspection (text, formulas, stale terms, literal 'None')
    print("\n--- PHASE 1: STATIC OPENPYXL SCAN ---")
    for t in TARGETS:
        p = ROOT / "companies" / t / "template_v6" / f"{t}-FullTemplate.xlsx"
        wb = openpyxl.load_workbook(p, data_only=False)
        print(f"\nChecking {t} ({p.name}, size: {p.stat().st_size} bytes):")
        
        # Check sheet count and names
        print(f"  Sheets ({len(wb.sheetnames)}): {wb.sheetnames}")
        
        stale_hits = []
        none_hits = []
        ref_hits = []
        
        for sname in wb.sheetnames:
            ws = wb[sname]
            for r in range(1, ws.max_row + 1):
                for c in range(1, ws.max_column + 1):
                    cell = ws.cell(r, c)
                    val = cell.value
                    if val is not None:
                        sval = str(val).lower()
                        col_let = openpyxl.utils.get_column_letter(c)
                        addr = f"{sname}!{col_let}{r}"
                        
                        # Stale terms
                        for term in STALE_TERMS:
                            if term in sval:
                                stale_hits.append((addr, term, val))
                        
                        # Literal "None" string
                        if sval == "none":
                            none_hits.append((addr, val))
                        
                        # #REF! in formula
                        if str(val).startswith("=") and "#REF!" in str(val):
                            ref_hits.append((addr, val))
                            
        print(f"  Stale term hits: {len(stale_hits)}")
        for addr, term, val in stale_hits:
            print(f"    STALE [{addr}] '{term}': {val}")
            
        print(f"  Literal 'None' string hits: {len(none_hits)}")
        for addr, val in none_hits[:10]:
            print(f"    LITERAL_NONE [{addr}]: {val}")
        if len(none_hits) > 10:
            print(f"    ... and {len(none_hits) - 10} more")
            
        print(f"  #REF! in formulas: {len(ref_hits)}")
        for addr, val in ref_hits:
            print(f"    REF_ERROR [{addr}]: {val}")

    # 2. Excel COM recalculation and live value inspection
    print("\n--- PHASE 2: EXCEL COM 16.0 RECALCULATION & ERROR SCAN ---")
    excel = win32.gencache.EnsureDispatch("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    
    try:
        for t in TARGETS:
            p = (ROOT / "companies" / t / "template_v6" / f"{t}-FullTemplate.xlsx").resolve()
            wb = excel.Workbooks.Open(str(p), UpdateLinks=0, ReadOnly=True)
            excel.CalculateFullRebuild()
            print(f"\n================ COM AUDIT: {t} ================")
            
            total_com_errs = 0
            for ws_idx in range(1, wb.Worksheets.Count + 1):
                ws = wb.Worksheets(ws_idx)
                sname = ws.Name
                used = ws.UsedRange
                start_row = used.Row
                start_col = used.Column
                vals = used.Value
                formulas = used.Formula
                
                if vals is None:
                    continue
                nr = len(vals)
                nc = len(vals[0]) if nr > 0 else 0
                
                sheet_errs = []
                for r in range(nr):
                    for c in range(nc):
                        v = vals[r][c]
                        err_name = None
                        if isinstance(v, int) and v in COM_ERRORS:
                            err_name = COM_ERRORS[v]
                        elif isinstance(v, str) and v in COM_ERRORS.values():
                            err_name = v
                        if err_name:
                            addr = f"{openpyxl.utils.get_column_letter(start_col + c)}{start_row + r}"
                            f = formulas[r][c] if formulas else ""
                            sheet_errs.append((addr, err_name, f))
                            
                if sheet_errs:
                    print(f"  Sheet [{sname}] has {len(sheet_errs)} errors:")
                    for addr, err, f in sheet_errs:
                        print(f"     {addr}: {err} (formula: {f})")
                    total_com_errs += len(sheet_errs)
                    
            print(f"  Total COM errors for {t}: {total_com_errs}")
            
            # Key checks
            dcf_ws = wb.Worksheets("DCF 1st way")
            model_ws = wb.Worksheets("Model ->")
            e42 = dcf_ws.Range("E42").Value
            c9 = model_ws.Range("C9").Value
            print(f"  'DCF 1st way'!E42 = {e42}")
            print(f"  'Model ->'!C9     = {c9}")
            print(f"  Difference: {abs(float(e42) - float(c9)):.6f}")
            
            # Check sensitivity base cell
            if t != "EDELWEISS":
                j27 = dcf_ws.Range("J27").Value
                e36 = dcf_ws.Range("E36").Value
                print(f"  Sensitivity J27 = {j27}, E36 = {e36}, Diff = {abs(float(j27) - float(e36)):.6f}")
                
            # Check BS reconciliation
            bs_ws = wb.Worksheets("BS")
            cols = ["C", "D", "E", "F", "G", "H", "I", "J", "K", "L", "M", "N"]
            bs_diffs = []
            for c_let in cols:
                liab = bs_ws.Range(f"{c_let}16").Value
                assets = bs_ws.Range(f"{c_let}38").Value
                if liab is not None and assets is not None:
                    d = abs(float(liab) - float(assets))
                    if d > 0.01:
                        bs_diffs.append((c_let, liab, assets, d))
            if bs_diffs:
                print(f"  BS MISMATCHES ({len(bs_diffs)}): {bs_diffs}")
            else:
                print("  BS: 100% Reconciled (Assets == Liabilities across all columns)")
                
            wb.Close(False)
    finally:
        excel.Quit()
        
    print("\n" + "="*60)
    print("COMPREHENSIVE DEEP AUDIT COMPLETED")
    print("="*60)

if __name__ == "__main__":
    run_comprehensive_audit()
