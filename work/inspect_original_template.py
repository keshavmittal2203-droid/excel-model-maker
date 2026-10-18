import openpyxl

wb = openpyxl.load_workbook('originals/Marksans Keshav.xlsx', data_only=False)
ws = wb['DCF 1st way']
print("="*60)
print("ORIGINAL TEMPLATE (originals/Marksans Keshav.xlsx): DCF 1st way")
print("="*60)
for r in range(1, 20):
    vals = [f"{c}{r}={ws[f'{c}{r}'].value}" for c in ['D', 'E', 'F', 'G', 'H', 'I', 'J', 'K'] if ws[f'{c}{r}'].value is not None]
    if vals:
        print(f"Row {r:2d}: " + ' | '.join(vals))
for r in range(20, 45):
    vals = [f"{c}{r}={ws[f'{c}{r}'].value}" for c in ['D', 'E'] if ws[f'{c}{r}'].value is not None]
    if vals:
        print(f"Row {r:2d}: " + ' | '.join(vals))

wacc = wb['WACC-DCF']
print("\nORIGINAL TEMPLATE: WACC-DCF Key Cells:")
for cell_id in ['D18', 'E22', 'E23', 'E24', 'E30', 'E31', 'E32', 'J22', 'J23', 'J24', 'J25', 'J32', 'J39']:
    c = wacc[cell_id]
    cmt = f" [Comment: {c.comment.text[:30]}...]" if c.comment else ""
    print(f"  {cell_id}: {c.value}{cmt}")
